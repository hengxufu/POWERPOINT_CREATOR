import tempfile
import unittest
import zipfile
from pathlib import Path
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from PIL import Image
from merge_decks import merge, xml, relpath, target


class MergeTests(unittest.TestCase):
    def test_graph_chart_notes_and_order(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            Image.new('RGB', (20, 20), 'blue').save(root/'image.png')
            for name in ('A', 'B'):
                prs = Presentation()
                slide = prs.slides.add_slide(prs.slide_layouts[6])
                slide.shapes.add_textbox(0, 0, 1000000, 1000000).text = name
                slide.notes_slide.notes_text_frame.text = name+'讲稿'
                chart = ChartData(); chart.categories = ['x']; chart.add_series(name, [3])
                slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, 0, 1000000, 3000000, 2000000, chart)
                slide.shapes.add_picture(str(root/'image.png'), 0, 0)
                prs.save(root/(name+'.pptx'))
            output = root/'merged.pptx'
            merge({'sources': [{'path':str(root/'B.pptx')}, {'path':str(root/'A.pptx')}]}, output)
            result = Presentation(output)
            self.assertEqual([s.shapes[0].text for s in result.slides], ['B', 'A'])
            self.assertEqual([s.notes_slide.notes_text_frame.text for s in result.slides], ['B讲稿', 'A讲稿'])
            self.assertEqual(result.slides[1].shapes[1].chart.series[0].values[0], 3)
            self.assertEqual(result.slides[0].shapes[2].image.blob, (root/'image.png').read_bytes())
            with zipfile.ZipFile(output) as z:
                for name in z.namelist():
                    if not name.endswith('.rels') or name == '_rels/.rels':
                        continue
                    owner = name.replace('/_rels/', '/')[:-5]
                    for rel in xml(z.read(name)):
                        if rel.get('TargetMode') != 'External':
                            self.assertIn(target(owner, rel.get('Target')), z.namelist())

    def test_rejects_different_canvas_without_output(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            for number in (1, 2):
                prs = Presentation(); prs.slide_width = 9000000 + number
                prs.slides.add_slide(prs.slide_layouts[6]); prs.save(root/f'{number}.pptx')
            with self.assertRaises(ValueError):
                merge({'sources':[{'path':str(root/'1.pptx')}, {'path':str(root/'2.pptx')}]}, root/'out.pptx')
            self.assertFalse((root/'out.pptx').exists())


if __name__ == '__main__':
    unittest.main()
