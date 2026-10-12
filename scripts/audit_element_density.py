"""Count slide objects and propose individual web assets for pages below 40."""
import argparse
import hashlib
import json
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

MIN_ELEMENTS = 40


def count_objects(shapes, canvas, transform=(1, 1, 0, 0)):
    sx, sy, ox, oy = transform
    items = []
    for shape in shapes:
        metadata = shape._element.xpath('.//p:cNvPr')
        if metadata and metadata[0].get('hidden') in ('1', 'true'):
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            xfrm = shape._element.grpSpPr.xfrm
            child_sx = shape.width / xfrm.chExt.cx if xfrm.chExt.cx else 1
            child_sy = shape.height / xfrm.chExt.cy if xfrm.chExt.cy else 1
            items.extend(count_objects(shape.shapes, canvas,
                         (sx * child_sx, sy * child_sy,
                          ox + sx * (shape.left - child_sx * xfrm.chOff.x),
                          oy + sy * (shape.top - child_sy * xfrm.chOff.y))))
            continue
        if shape.shape_type in (MSO_SHAPE_TYPE.TEXT_BOX, MSO_SHAPE_TYPE.PLACEHOLDER):
            if not (shape.has_text_frame and shape.text.strip()) and not shape.has_table and not shape.has_chart and not hasattr(shape, 'image'):
                continue
        x, y = ox + sx * shape.left, oy + sy * shape.top
        width, height = sx * shape.width, sy * shape.height
        if width < 0 or height < 0 or width + height <= 0:
            continue
        if x < -1 or y < -1 or x + width > canvas[0] + 1 or y + height > canvas[1] + 1:
            continue
        items.append({'name': shape.name, 'type': str(shape.shape_type),
                      'picture': hasattr(shape, 'image')})
    return items


def audit(pptx, minimum=MIN_ELEMENTS, plan=None):
    if type(minimum) is not int or minimum < 1:
        raise ValueError('Minimum must be a positive integer')
    pptx = Path(pptx)
    prs = Presentation(pptx)
    plan = plan or {}
    pages = []
    used_ids = {request.get('request_id') for slide in plan.get('slides', []) for request in slide.get('asset_requests', [])}
    for number, slide in enumerate(prs.slides, 1):
        objects = count_objects(slide.shapes, (prs.slide_width, prs.slide_height))
        deficit = max(0, minimum - len(objects))
        planned = next((item for item in plan.get('slides', []) if item.get('slide_number') == number), {})
        title = planned.get('title') or (slide.shapes.title.text if slide.shapes.title else f'第{number}页')
        slide_id = planned.get('slide_id', f'slide-{number:02d}')
        requests = []
        for index in range(1, deficit + 1):
            identifier = f'density-s{number:03d}-e{index:03d}'
            while identifier in used_ids:
                identifier += '-v'
            used_ids.add(identifier)
            requests.append({'request_id': identifier, 'asset_id': identifier, 'object_id': identifier,
                             'kind': 'icon' if index % 2 else 'ornament', 'decision': 'generate_web',
                             'purpose': f'补足第{number}页独立元素；布局前核对与本页论证的相关性和留白',
                             'single_subject_prompt': f'围绕页面主题“{title}”，生成第{index}个独立的微型科研图标或装饰元素。'
                               '只有一个完整主体，透明背景，无文字、数字、数据、公式、页面布局或拼贴。'
                               '使用本稿设计系统；与已生成元素区分。仅概念与装饰，不代表实测结果。'})
        pages.append({'slide_number': number, 'slide_id': slide_id, 'title': title,
                      'element_count': len(objects), 'picture_count': sum(item['picture'] for item in objects),
                      'minimum': minimum, 'deficit': deficit, 'objects': objects,
                      'suggested_web_asset_requests': requests})
    return {'schema': 'ppt_element_density_v1', 'pptx_sha256': hashlib.sha256(pptx.read_bytes()).hexdigest(),
            'minimum_elements_per_slide': minimum, 'ready': bool(pages) and all(page['deficit'] == 0 for page in pages),
            'slides': pages, 'total_deficit': sum(page['deficit'] for page in pages),
            'request_status': 'proposals_not_submitted',
            'limits': 'Counts independent slide objects, recursively expands groups, excludes hidden, empty text and out-of-canvas objects; visual inspection is still required.'}


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--pptx', type=Path, required=True)
    parser.add_argument('--plan', type=Path)
    parser.add_argument('--out', type=Path, required=True)
    args = parser.parse_args()
    if args.out.resolve() in {args.pptx.resolve(), args.plan.resolve() if args.plan else args.pptx.resolve()}:
        parser.error('Report must not overwrite inputs')
    plan = json.loads(args.plan.read_text(encoding='utf-8-sig')) if args.plan else None
    result = audit(args.pptx, plan=plan)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'ready': result['ready'], 'total_deficit': result['total_deficit']}, ensure_ascii=False))
    raise SystemExit(0 if result['ready'] else 1)
