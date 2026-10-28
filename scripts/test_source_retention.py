import tempfile
import unittest
from pathlib import Path
from pptx import Presentation
from audit_source_retention import audit


class RetentionTests(unittest.TestCase):
    def test_notes_tags_and_required_text(self):
        with tempfile.TemporaryDirectory() as folder:
            deck = Path(folder) / 'deck.pptx'
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            box = slide.shapes.add_textbox(0, 0, 1000000, 1000000)
            paragraph = box.text_frame.paragraphs[0]
            paragraph.add_run().text = '{{TI'
            paragraph.add_run().text = 'TLE}}'
            slide.notes_slide.notes_text_frame.text = '真实讲稿'
            prs.save(deck)
            result = audit(deck, {'slides': [{'speaker_notes': '其他讲稿', 'required_visible_text': ['结果']} ]})
            self.assertEqual({v['code'] for v in result['findings']},
                             {'unresolved_template_tags', 'speaker_notes_mismatch', 'required_text_missing'})
            box.text = '结果'
            prs.save(deck)
            self.assertTrue(audit(deck, {'slides': [{'speaker_notes': '真实讲稿', 'required_visible_text': ['结果']} ]})['ready'])


if __name__ == '__main__':
    unittest.main()
