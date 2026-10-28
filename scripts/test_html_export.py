"""Opt-in real-browser regression; PPT_HTML_INTEGRATION=1 enables it."""
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from pptx import Presentation


@unittest.skipUnless(os.environ.get('PPT_HTML_INTEGRATION') == '1', 'Real browser integration opt-in')
class HtmlTests(unittest.TestCase):
    def test_native_text_notes_and_missing_image(self):
        runner = Path(__file__).with_name('local_engine.py')
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            html = root/'deck.html'; output = root/'deck.pptx'
            html.write_text('<meta charset="utf-8"><style>.slide{width:960px;height:540px;font:32px Arial}</style>'
                            '<section class="slide" data-pptx-notes="讲稿A"><h1>中文 Title</h1><p>Value 42</p></section>'
                            '<section class="slide" data-pptx-notes="讲稿B"><h1>Second</h1></section>', 'utf-8')
            command = [sys.executable, '-X', 'utf8', str(runner), 'html', '--', str(html), '--output', str(output), '--inject']
            result = subprocess.run(command, capture_output=True, text=True, encoding='utf-8', timeout=120)
            self.assertEqual(result.returncode, 0, result.stdout+result.stderr)
            prs = Presentation(output)
            self.assertEqual(len(prs.slides), 2)
            self.assertIn('中文 Title', '\n'.join(s.text for s in prs.slides[0].shapes if s.has_text_frame))
            self.assertIn('讲稿A', prs.slides[0].notes_slide.notes_text_frame.text)
            original = output.read_bytes()
            duplicate = subprocess.run(command, capture_output=True, timeout=120)
            self.assertNotEqual(duplicate.returncode, 0)
            self.assertEqual(output.read_bytes(), original)
            html.write_text('<div class="slide"><img src="missing.png"></div>', 'utf-8')
            command[command.index(str(output))] = str(root/'broken.pptx')
            broken = subprocess.run(command, capture_output=True, timeout=120)
            self.assertNotEqual(broken.returncode, 0)
            self.assertFalse((root/'broken.pptx').exists())


if __name__ == '__main__':
    unittest.main()
