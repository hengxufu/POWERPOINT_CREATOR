"""Offline capability verification; sample values and planning are synthetic."""
import copy
import json
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from types import SimpleNamespace
from pptx import Presentation
from native_data import build, validate, TYPES
from plan_workflow import packet, jobs
from typography_check import audit


def example():
    slides = []
    for kind in TYPES:
        item = {'id': kind + '-editable', 'type': 'chart', 'chart_type': kind, 'box': [.7, 2, 8, 4.5],
                'source': '合成示例，非实测数据', 'unit': '示例得分', 'categories': ['条件A', '条件B', '条件C'],
                'series': [{'name': '合成序列', 'values': [24, 48, 65]}]}
        if kind == 'scatter':
            item['series'] = [{'name': '合成序列', 'points': [[1, 24], [2, 48], [3, 65]]}]
            item['x_unit'] = '示例时间'
        slides.append({'title': f'原生数据对象｜{kind}（合成示例）', 'takeaway': '保留可编辑图表和嵌入工作簿；不代表科学结论', 'items': [item]})
    slides.append({'title': '原生证据表｜合成示例', 'takeaway': '文字、列宽、单元格均可修改', 'items': [
        {'id': 'table-editable', 'type': 'table', 'box': [.7, 2.1, 11.8, 3.6], 'source': '能力演示，不是真实项目材料',
         'unit': 'not_applicable', 'column_weights': [1, 2, 2], 'rows': [
             ['比较维度', '方案A（示例）', '方案B（示例）'], ['条件', '相同测试条件', '相同测试条件'],
             ['结果', '24（合成数据）', '48（合成数据）'], ['限制', '未提供真实验证', '未提供真实验证']]}]})
    return {'profile': 'research', 'slides': slides}


def main():
    checks = []
    with tempfile.TemporaryDirectory() as temporary:
        root = Path(temporary)
        data = example()
        deck = root / 'data.pptx'
        build(data, deck)
        prs = Presentation(deck)
        assert len(prs.slides) == 8
        charts = [shape.chart for slide in prs.slides for shape in slide.shapes if shape.has_chart]
        assert len(charts) == 7 and all(chart.part.chart_workbook.xlsx_part.blob for chart in charts)
        checks.append('seven chart types retain embedded workbooks')
        for chart in charts:
            for node in chart._chartSpace.xpath('.//c:axId | .//c:crossAx'):
                assert 0 <= int(node.get('val')) < 2 ** 32
        assert charts[-1]._chartSpace.xpath('.//c:scatterStyle')[0].get('val') == 'marker'
        checks.append('unsigned axis references and unconnected scatter semantics preserved')
        table = next(shape.table for shape in prs.slides[-1].shapes if shape.has_table)
        assert table.cell(2, 1).text == '24（合成数据）'
        checks.append('native table cells match input')
        assert audit(deck)['passed_explicit_font_check']
        checks.append('CJK and Latin typography passes')
        revised = copy.deepcopy(data)
        revised['slides'][0]['items'][0]['series'][0]['values'][0] = 39
        revised['slides'][-1]['items'][0]['rows'][2][1] = '39（修改后示例）'
        rebuilt = root / 'rebuilt.pptx'
        build(revised, rebuilt)
        changed = Presentation(rebuilt)
        chart = next(shape.chart for shape in changed.slides[0].shapes if shape.has_chart)
        assert chart.series[0].values[0] == 39
        with zipfile.ZipFile(__import__('io').BytesIO(chart.part.chart_workbook.xlsx_part.blob)) as workbook:
            assert b'39' in workbook.read('xl/worksheets/sheet1.xml')
        assert next(shape.table for shape in changed.slides[-1].shapes if shape.has_table).cell(2, 1).text == '39（修改后示例）'
        checks.append('chart value embedded workbook and table update on rebuild')
        for modifier in ('length', 'nan', 'table', 'source'):
            bad = copy.deepcopy(data)
            if modifier == 'length': bad['slides'][0]['items'][0]['series'][0]['values'] = [1]
            if modifier == 'nan': bad['slides'][0]['items'][0]['series'][0]['values'][0] = float('nan')
            if modifier == 'table': bad['slides'][-1]['items'][0]['rows'][1] = ['short']
            if modifier == 'source': bad['slides'][0]['items'][0]['source'] = ''
            try: validate(bad)
            except ValueError: pass
            else: raise AssertionError(modifier)
            checks.append('invalid ' + modifier + ' rejected')
        brief = root / 'brief.md'
        brief.write_text('合成任务：设计一个产品概念与准确字标；无真实研究结果。', encoding='utf-8')
        packet(SimpleNamespace(brief=brief, profile='challenge', out=root / 'packet'))
        assert 'prepared' not in (root / 'packet/gpt6-task-packet.md').read_text(encoding='utf-8')
        checks.append('scene-specific planning packet created without claiming web response')
        plan = root / 'plan.json'
        plan.write_text(json.dumps({'slides': [{'slide_number': 1, 'slide_id': 's01', 'asset_requests': [
            {'request_id': 'icon-one', 'object_id': 'icon', 'kind': 'icon', 'single_subject_prompt': '一个完整传感器图标'},
            {'request_id': 'wordmark-one', 'object_id': 'wordmark', 'kind': 'art_text', 'single_subject_prompt': '独立字标', 'exact_text': '融合感知'}]}]}), encoding='utf-8')
        jobs(SimpleNamespace(plan=plan, profile='challenge', out=root / 'queue'))
        queue = json.loads((root / 'queue/queue.json').read_text(encoding='utf-8'))
        assert len(queue['jobs']) == 2
        assert all(json.loads((root / 'queue' / entry['job_dir'] / 'job.json').read_text(encoding='utf-8'))['status'] == 'prepared' for entry in queue['jobs'])
        checks.append('one task per icon or wordmark; no fabricated submission')
    for check in checks: print('PASS:', check)
    print(f'{len(checks)} checks passed')


if __name__ == '__main__':
    main()
