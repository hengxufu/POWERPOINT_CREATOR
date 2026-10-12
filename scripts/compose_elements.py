#!/usr/bin/env python3
"""Insert reviewed single-element image jobs into a PPTX as separate picture objects."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches


ALLOWED_KINDS = {"ornament", "concept", "art_text", "icon", "cutout"}


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def require(condition: bool, message: str) -> None:
    if not condition:
        raise ValueError(message)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pptx", required=True, type=Path)
    parser.add_argument("--manifest", required=True, type=Path)
    parser.add_argument("--out", required=True, type=Path)
    parser.add_argument("--report", required=True, type=Path)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    source = args.pptx.resolve()
    manifest_path = args.manifest.resolve()
    output = args.out.resolve()
    report_path = args.report.resolve()
    require(source.is_file(), f"PPTX not found: {source}")
    require(manifest_path.is_file(), f"Manifest not found: {manifest_path}")
    require(not output.exists(), f"Refusing to overwrite output: {output}")
    require(not report_path.exists(), f"Refusing to overwrite report: {report_path}")
    require(output != report_path, "PPTX output and report must use different paths")
    require(output != manifest_path and report_path != manifest_path, "Outputs must not overwrite the manifest")

    manifest = load_json(manifest_path)
    require(manifest.get("schema") == "ppt_element_manifest_v1", "Unsupported manifest schema")
    max_fraction = float(manifest.get("max_generated_asset_area_fraction", 0.55))
    require(0 < max_fraction <= 0.55, "max_generated_asset_area_fraction must be in (0, 0.55]")

    prs = Presentation(str(source))
    slide_width = prs.slide_width / 914400
    slide_height = prs.slide_height / 914400
    page_area = slide_width * slide_height
    inserted = []
    used_names: set[str] = set()

    for index, element in enumerate(manifest.get("elements", []), start=1):
        kind = element.get("kind")
        require(kind in ALLOWED_KINDS, f"Element {index}: forbidden kind {kind!r}")
        slide_number = element.get("slide_number", 0)
        require(type(slide_number) is int, f"Element {index}: slide_number must be an integer")
        require(1 <= slide_number <= len(prs.slides), f"Element {index}: invalid slide_number")

        job_dir = Path(element["job_dir"])
        if not job_dir.is_absolute():
            job_dir = (manifest_path.parent / job_dir).resolve()
        job = load_json(job_dir / "job.json")
        placement = load_json(job_dir / "placement.json")
        require(job.get("status") in {"accepted", "integrated"}, f"Element {index}: job is not accepted")
        spec = job.get("spec", {})
        require(spec.get("kind") == kind, f"Element {index}: kind mismatch")

        slide_id = str(spec.get("slide_id", "")).strip()
        object_id = str(spec.get("object_id", "")).strip()
        require(slide_id and object_id, f"Element {index}: missing slide_id/object_id")
        shape_name = f"{slide_id}/{object_id}"
        require(shape_name not in used_names, f"Element {index}: duplicate object name {shape_name}")
        require(not any(shape.name == shape_name for page in prs.slides for shape in page.shapes),
                f"Element {index}: shape already exists in deck")
        used_names.add(shape_name)
        require(placement.get("schema") == "single_element_placement_v1", f"Element {index}: unsupported placement schema")
        for field, expected in (("slide_id", slide_id), ("object_id", object_id), ("shape_name", shape_name)):
            require(placement.get(field) == expected, f"Element {index}: placement {field} mismatch")

        box = placement.get("image_box_inches")
        require(isinstance(box, list) and len(box) == 4, f"Element {index}: missing image_box_inches")
        x, y, width, height = map(float, box)
        require(all(math.isfinite(value) for value in (x, y, width, height)), f"Element {index}: non-finite placement")
        require(x >= 0 and y >= 0 and width > 0 and height > 0, f"Element {index}: invalid placement")
        require(x + width <= slide_width + 1e-6, f"Element {index}: placement exceeds slide width")
        require(y + height <= slide_height + 1e-6, f"Element {index}: placement exceeds slide height")
        require((width * height) / page_area <= max_fraction + 1e-9, f"Element {index}: generated asset exceeds area limit")

        image_value = job.get("asset_file")
        require(bool(image_value), f"Element {index}: missing image path")
        image_path = Path(image_value)
        if not image_path.is_absolute():
            image_path = (job_dir / image_path).resolve()
        require(image_path.is_file(), f"Element {index}: image not found: {image_path}")
        placement_path = Path(placement.get("asset_file", ""))
        if not placement_path.is_absolute():
            placement_path = job_dir / placement_path
        require(placement_path.resolve() == image_path.resolve(), f"Element {index}: placement image differs from accepted job")
        actual_hash = sha256(image_path)
        expected_hash = job.get("sha256")
        require(bool(expected_hash), f"Element {index}: missing expected image hash")
        require(placement.get("asset_sha256") == expected_hash, f"Element {index}: placement hash differs from accepted job")
        require(actual_hash.lower() == str(expected_hash).lower(), f"Element {index}: image hash mismatch")
        with Image.open(image_path) as image:
            pixel_width, pixel_height = image.size
        require(math.isclose(width / height, pixel_width / pixel_height, rel_tol=1e-4),
                f"Element {index}: placement would stretch the accepted image; regenerate placement")

        slide = prs.slides[slide_number - 1]
        picture = slide.shapes.add_picture(
            str(image_path), Inches(x), Inches(y), width=Inches(width), height=Inches(height)
        )
        picture.name = shape_name
        c_nv_pr = picture._element.nvPicPr.cNvPr
        c_nv_pr.set("title", f"{kind}: {object_id}")
        c_nv_pr.set("descr", f"single-element asset; sha256={actual_hash}; source={image_path}")
        inserted.append({
            "slide_number": slide_number,
            "shape_name": shape_name,
            "kind": kind,
            "image_path": str(image_path),
            "image_sha256": actual_hash,
            "placement_in": {"x": x, "y": y, "width": width, "height": height},
        })

    output.parent.mkdir(parents=True, exist_ok=True)
    report_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
    report = {
        "schema": "ppt_element_composition_report_v1",
        "source_pptx": str(source),
        "source_sha256": sha256(source),
        "output_pptx": str(output),
        "output_sha256": sha256(output),
        "manifest": str(manifest_path),
        "manifest_sha256": sha256(manifest_path),
        "inserted_count": len(inserted),
        "delivery_status": "draft_requires_delivery_gate",
        "warning": "Composition alone does not establish that planned assets, content or visual review are complete.",
        "inserted": inserted,
        "whole_slide_generation_used": False,
    }
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(output), "inserted_count": len(inserted)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
