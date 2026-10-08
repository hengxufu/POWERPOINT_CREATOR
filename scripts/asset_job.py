"""Prepare and track single-element WEB generation. Does not drive a browser."""
import argparse
import datetime
import hashlib
import json
import math
import re
import shutil
from pathlib import Path
from PIL import Image

def now():
    return datetime.datetime.now(datetime.timezone.utc).isoformat()

def write(path, value):
    temporary = path.with_suffix(path.suffix+'.tmp')
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')
    temporary.replace(path)

def read(job):
    return json.loads((job/'job.json').read_text(encoding='utf-8'))

def prompt(spec):
    if not re.fullmatch(r'[a-zA-Z0-9][a-zA-Z0-9_-]{0,63}', spec.get('asset_id','')):
        raise ValueError('asset_id must be a short filename-safe identifier')
    for key in ('element','slide_id','object_id'):
        if not isinstance(spec.get(key),str) or not spec[key].strip():
            raise ValueError('Missing '+key)
    if spec.get('kind') not in ('ornament','concept','art_text'):
        raise ValueError('kind must be ornament, concept or art_text')
    if spec.get('kind')=='art_text' and not spec.get('exact_text'):
        raise ValueError('art_text requires exact_text')
    if not isinstance(spec.get('transparent',True),bool):
        raise ValueError('transparent must be boolean')
    for key in ('min_width','min_height'):
        value=spec.get(key,512)
        if type(value) is not int or not 1 <= value <= 16384:
            raise ValueError(key+' must be an integer 1..16384')
    lines=['请实际生成一张图片，只包含一个独立元素，供可编辑 PPT 作为局部素材使用。',
           '主体：'+spec['element'],
           '风格：'+spec.get('style','克制、精致，适合中国科研汇报'),
           '颜色：'+spec.get('palette','深青色与少量浅青高光'),
           '构图比例：'+spec.get('aspect','1:1')+'；主体完整，四周保留约 10% 安全留白。',
           '真实 alpha 透明背景，不要棋盘格模拟透明。' if spec.get('transparent',True) else '背景：'+spec.get('background','纯白'),
           f"目标尺寸至少 {spec.get('min_width',512)} × {spec.get('min_height',512)} 像素；这是期望，不是 API 参数。",
           '禁止整页幻灯片、标题区、正文、页脚、目录、拼图、九宫格、多个变体、边框、水印或品牌标识。']
    if spec['kind']=='art_text':
        lines.append('唯一允许的文字是：'+json.dumps(spec['exact_text'],ensure_ascii=False)+'。保持逐字准确；不要附加任何其他文字。')
    else:
        lines.append('图片内禁止出现任何文字、数字或标注；需要的说明会在 PPT 中以原生文本加入。')
    if spec['kind']=='concept':
        lines.append('这是概念性插画，不是实验图或科学证据；不要捏造数据、测量值或未经说明的机制关系。')
    context=spec.get('visual_context',{})
    if not isinstance(context,dict): raise ValueError('visual_context must be an object')
    for key,label in [('slide_background','使用页背景'),('lighting','光照'),('material','材质'),('placement','摆放位置'),('avoid','额外避免')]:
        if context.get(key): lines.append(label+'：'+str(context[key]))
    if context: lines.append('这些是局部素材适配条件，不要把背景色、文字区或幻灯片本身画进图中。')
    return '\n'.join(lines)

def prepare(args):
    spec=json.loads(Path(args.spec).read_text(encoding='utf-8'))
    result=prompt(spec)
    job=Path(args.job)
    job.mkdir(parents=True,exist_ok=False)
    (job/'prompt.txt').write_text(result,encoding='utf-8')
    write(job/'job.json', {'version':1,'transport':'chatgpt_web','requested_engine':'image-2.5',
                          'verified_engine':None,'status':'prepared','spec':spec,
                          'history':[{'at':now(),'status':'prepared'}]})

def event(args):
    job=Path(args.job); data=read(job)
    statuses={'blocked_browser','blocked_model','blocked_download','submission_unknown','submitted','generated','integrated'}
    if args.status not in statuses: raise ValueError('Unsupported status')
    if args.status=='submitted':
        if any(e['status'] in ('submitted','submission_unknown') for e in data['history']):
            raise ValueError('Submission already recorded; inspect the existing web session')
        if not args.session_url or not re.fullmatch(r'https://chatgpt\.com/c/[A-Za-z0-9_-]+',args.session_url):
            raise ValueError('submitted requires an observed ChatGPT conversation URL')
        data['session_url']=args.session_url
        if args.engine and args.engine!='unverified':
            if not args.engine_evidence: raise ValueError('Engine needs visible UI evidence')
            data['verified_engine']=args.engine
            data['engine_evidence']=args.engine_evidence
    if args.status=='integrated' and data['status']!='accepted':
        raise ValueError('Integration needs accepted visual review first')
    if args.status=='integrated':
        if not getattr(args,'pptx',None): raise ValueError('Actual PPTX required to verify integration')
        from pptx import Presentation
        deck=Path(args.pptx)
        target=data['spec']['slide_id']+'/'+data['spec']['object_id']
        found=[shape for slide in Presentation(deck).slides for shape in slide.shapes if shape.name==target]
        if len(found)!=1 or not hasattr(found[0],'image'): raise ValueError('Expected one native picture named '+target)
        if hashlib.sha256(found[0].image.blob).hexdigest()!=data['sha256']: raise ValueError('PPTX picture does not match accepted asset')
        data['integration']={'pptx':str(deck.resolve()),'pptx_sha256':hashlib.sha256(deck.read_bytes()).hexdigest(),'shape_name':target,'at':now()}
    if args.status in ('generated','blocked_download') and not any(e['status']=='submitted' for e in data['history']):
        raise ValueError('Generation/download status requires actual recorded submission')
    data['status']=args.status
    data['history'].append({'at':now(),'status':args.status,'note':args.note})
    write(job/'job.json',data)

def receive(args):
    job=Path(args.job); data=read(job)
    if data['status'] not in ('submitted','generated','blocked_download'):
        raise ValueError('Record actual web submission before receiving its download')
    if not any(e['status']=='submitted' for e in data['history']):
        raise ValueError('Missing actual submission record')
    path=Path(args.file)
    with Image.open(path) as im:
        im.load()
        fmt=im.format
        if fmt not in ('PNG','JPEG','WEBP'): raise ValueError('Expected PNG, JPEG or WEBP')
        alpha_channel=im.convert('RGBA').getchannel('A')
        alpha=alpha_channel.getextrema()
        bbox=alpha_channel.point(lambda x:255 if x>=16 else 0).getbbox()
        checks={'width':im.width,'height':im.height,'format':fmt,
                'has_transparent_pixels':alpha[0]<255,
                'has_visible_pixels':alpha[1]>0,
                'subject_bbox_px':list(bbox) if bbox else None,
                'subject_box_area_fraction':round((bbox[2]-bbox[0])*(bbox[3]-bbox[1])/(im.width*im.height),4) if bbox else 0}
    spec=data['spec']
    checks['size_ok']=checks['width']>=spec.get('min_width',512) and checks['height']>=spec.get('min_height',512)
    checks['alpha_ok']=not spec.get('transparent',True) or checks['has_transparent_pixels']
    checks['warnings']=[]
    if checks['subject_box_area_fraction']<.15: checks['warnings'].append('主体包围盒占比小，放置后可能显得过小；人工检查，不自动裁切')
    if bbox and (bbox[0]==0 or bbox[1]==0 or bbox[2]==checks['width'] or bbox[3]==checks['height']): checks['warnings'].append('主体接触图片边缘，请确认完整性与安全留白')
    suffix={'PNG':'.png','JPEG':'.jpg','WEBP':'.webp'}[fmt]
    dest=job/(spec['asset_id']+'-original'+suffix)
    if dest.exists(): raise ValueError('Original asset already exists; use a new version job')
    shutil.copyfile(path,dest)
    data['asset_file']=dest.name
    data['sha256']=hashlib.sha256(dest.read_bytes()).hexdigest()
    data['technical_checks']=checks
    data['status']='received' if checks['size_ok'] and checks['alpha_ok'] and checks['has_visible_pixels'] else 'technical_failed'
    data['history'].append({'at':now(),'status':data['status']})
    write(job/'job.json',data)

def review(args):
    job=Path(args.job); data=read(job)
    if data['status'] not in ('received','rejected'): raise ValueError('Technical checks must pass before visual acceptance')
    if not args.note.strip(): raise ValueError('Actual visual review notes required')
    data['status']=args.verdict
    data['visual_review']={'note':args.note,'at':now()}
    data['history'].append({'at':now(),'status':args.verdict,'note':args.note})
    write(job/'job.json',data)

def placement(args):
    job=Path(args.job); data=read(job)
    if data['status'] not in ('accepted','integrated'): raise ValueError('Accepted review required before placement')
    asset=job/data['asset_file']
    if hashlib.sha256(asset.read_bytes()).hexdigest()!=data['sha256']: raise ValueError('Accepted asset modified; re-review a new version')
    x,y,w,h=args.box; cw,ch=args.canvas
    if not all(math.isfinite(n) for n in [x,y,w,h,cw,ch,args.min_ppi]): raise ValueError('Finite numbers required')
    if min(w,h,cw,ch,args.min_ppi)<=0 or min(x,y)<0 or x+w>cw or y+h>ch: raise ValueError('Target box must stay inside canvas')
    checks=data['technical_checks']; pw,ph=checks['width'],checks['height']
    scale=min(w/pw,h/ph); iw,ih=pw*scale,ph*scale
    ix=x+(w-iw)/2; iy=y+(h-ih)/2; ppi=pw/iw
    bbox=checks.get('subject_bbox_px')
    plan={'schema':'single_element_placement_v1','slide_id':data['spec']['slide_id'],'object_id':data['spec']['object_id'],
          'shape_name':data['spec']['slide_id']+'/'+data['spec']['object_id'],'asset_file':str(asset.resolve()),'asset_sha256':data['sha256'],
          'editable':False,'semantic_text':data['spec'].get('exact_text'),'target_box_inches':[x,y,w,h],
          'image_box_inches':[round(v,6) for v in [ix,iy,iw,ih]],'effective_ppi':round(ppi,1),'resolution_ok':ppi>=args.min_ppi,
          'visible_subject_box_inches':[round(v,6) for v in [ix+bbox[0]*scale,iy+bbox[1]*scale,(bbox[2]-bbox[0])*scale,(bbox[3]-bbox[1])*scale]] if bbox else None,
          'warnings':checks.get('warnings',[])+([] if ppi>=args.min_ppi else ['分辨率低于摆放要求，缩小或重新生成']),
          'image_editing':'none; original preserved; contain without stretching',
          'visual_review_required':['实际背景上的边缘','与文字和证据的间距','视觉大小和光照一致性']}
    write(job/'placement.json',plan); data['placement_file']='placement.json'; write(job/'job.json',data)

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    sub=parser.add_subparsers(dest='command',required=True)
    p=sub.add_parser('prepare'); p.add_argument('--spec',required=True); p.add_argument('--job',required=True); p.set_defaults(func=prepare)
    p=sub.add_parser('event'); p.add_argument('--job',required=True); p.add_argument('--status',required=True); p.add_argument('--note',required=True); p.add_argument('--session-url'); p.add_argument('--engine'); p.add_argument('--engine-evidence'); p.add_argument('--pptx'); p.set_defaults(func=event)
    p=sub.add_parser('receive'); p.add_argument('--job',required=True); p.add_argument('--file',required=True); p.set_defaults(func=receive)
    p=sub.add_parser('review'); p.add_argument('--job',required=True); p.add_argument('--verdict',choices=['accepted','rejected'],required=True); p.add_argument('--note',required=True); p.set_defaults(func=review)
    p=sub.add_parser('placement'); p.add_argument('--job',required=True); p.add_argument('--box',type=float,nargs=4,required=True); p.add_argument('--canvas',type=float,nargs=2,default=[13.333,7.5]); p.add_argument('--min-ppi',type=float,default=180); p.set_defaults(func=placement)
    args=parser.parse_args(); args.func(args)
    print(json.dumps(read(Path(args.job)),ensure_ascii=False,indent=2))

if __name__=='__main__': main()
