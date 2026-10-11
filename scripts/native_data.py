"""Build native editable charts and tables from a JSON specification.

This helper creates a new deck. Do not use it to round-trip an existing complex deck.
Data provenance and units are required; no inferred statistics or fabricated evidence.
"""
import argparse
import json
import math
from pathlib import Path
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.util import Inches, Pt
from typography_check import normalize

TYPES = {'column': XL_CHART_TYPE.COLUMN_CLUSTERED, 'bar': XL_CHART_TYPE.BAR_CLUSTERED,
         'line': XL_CHART_TYPE.LINE_MARKERS, 'area': XL_CHART_TYPE.AREA,
         'pie': XL_CHART_TYPE.PIE, 'doughnut': XL_CHART_TYPE.DOUGHNUT,
         'scatter': XL_CHART_TYPE.XY_SCATTER}
PALETTES = {'scholarship': ['991F28', 'C49A4A', '234364'],
            'research': ['174B70', '008E98', 'DA9253'],
            'challenge': ['0C5892', '00A9B0', 'DB8D3C'],
            'innovation': ['14776B', 'DE9744', '304C69']}


def textbox(slide, name, value, box, size=20, color='203446', bold=False):
    shape = slide.shapes.add_textbox(*(Inches(value) for value in box))
    shape.name = name
    shape.text_frame.word_wrap = True
    shape.text_frame.margin_left = shape.text_frame.margin_right = 0
    for index, line in enumerate(str(value).split('\n')):
        paragraph = shape.text_frame.paragraphs[0] if index == 0 else shape.text_frame.add_paragraph()
        run = paragraph.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)
    return shape


def validate(data):
    if not data.get('slides'):
        raise ValueError('At least one slide is required')
    for slide in data['slides']:
        ids = set()
        for item in slide.get('items', []):
            if not item.get('id') or item['id'] in ids:
                raise ValueError('Each item needs a unique id on its page')
            ids.add(item['id'])
            if not item.get('source') or not item.get('unit'):
                raise ValueError('Each data object requires source and unit (use not_applicable where appropriate)')
            x, y, width, height = item['box']
            if not all(math.isfinite(value) for value in item['box']) or min(x, y) < 0 or min(width, height) <= 0 or x + width > 13.3331 or y + height > 7.5001:
                raise ValueError('Data object must fit within 13.333 x 7.5 inches')
            if item['type'] == 'table':
                rows = item['rows']
                if not rows or not rows[0] or any(len(row) != len(rows[0]) for row in rows):
                    raise ValueError('Table rows must form a nonempty rectangle')
                if height / len(rows) < 0.3:
                    raise ValueError('Table is too dense; split across pages')
            elif item['type'] == 'chart':
                kind = item['chart_type']
                if kind not in TYPES or not item.get('series'):
                    raise ValueError('Unsupported chart type or empty series')
                for series in item['series']:
                    values = series['points'] if kind == 'scatter' else series['values']
                    if not values:
                        raise ValueError('Empty data series')
                    numbers = [coordinate for point in values for coordinate in point] if kind == 'scatter' else values
                    if any(type(value) not in (int, float) or not math.isfinite(value) for value in numbers):
                        raise ValueError('Values must be finite numbers; do not silently replace missing data')
                    if kind == 'scatter' and any(len(point) != 2 for point in values):
                        raise ValueError('Scatter points must be x,y pairs')
                    if kind != 'scatter' and len(values) != len(item['categories']):
                        raise ValueError('Category/value lengths differ')
                    if kind in ('pie', 'doughnut') and (min(values) < 0 or sum(values) <= 0):
                        raise ValueError('Composition charts require nonnegative values with positive total')
                if kind in ('pie', 'doughnut') and len(item['series']) != 1:
                    raise ValueError('Pie/doughnut accepts one series')
            else:
                raise ValueError('Only chart/table data objects are supported')


def add_chart(slide, item, palette):
    kind = item['chart_type']
    if kind == 'scatter':
        values = XyChartData()
        for series in item['series']:
            target = values.add_series(series['name'])
            for x, y in series['points']:
                target.add_data_point(x, y)
    else:
        values = CategoryChartData()
        values.categories = item['categories']
        for series in item['series']:
            values.add_series(series['name'], series['values'])
    shape = slide.shapes.add_chart(TYPES[kind], *(Inches(v) for v in item['box']), values)
    shape.name = item['id']
    chart = shape.chart
    chart.has_title = False
    chart.has_legend = len(item['series']) > 1 or kind in ('pie', 'doughnut')
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(14)
    for index, series in enumerate(chart.series):
        color = RGBColor.from_string(palette[index % len(palette)])
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = color
        if kind in ('line', 'scatter'):
            series.format.line.color.rgb = color
            series.format.line.width = Pt(2)
        if kind == 'scatter':
            series.format.line.fill.background()
            series.marker.style = XL_MARKER_STYLE.CIRCLE
            series.marker.size = 6
            series.marker.format.fill.solid()
            series.marker.format.fill.fore_color.rgb = color
        if kind in ('pie', 'doughnut'):
            for j, point in enumerate(series.points):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = RGBColor.from_string(palette[j % len(palette)])
    if kind not in ('pie', 'doughnut'):
        chart.value_axis.tick_labels.font.size = Pt(14)
        chart.category_axis.tick_labels.font.size = Pt(14)
        chart.value_axis.has_title = True
        chart.value_axis.axis_title.text_frame.text = item['unit']
        if kind in ('bar', 'column'):
            if item.get('value_min', 0) != 0:
                raise ValueError('Bar/column axes must include zero; custom truncation is refused')
            chart.value_axis.minimum_scale = min(0, min(v for s in item['series'] for v in s['values']))
        elif 'value_min' in item:
            chart.value_axis.minimum_scale = item['value_min']
        if kind == 'scatter':
            chart.category_axis.has_title = True
            chart.category_axis.axis_title.text_frame.text = item.get('x_unit', 'x')
    if kind == 'scatter':
        for element in chart._chartSpace.xpath('.//c:scatterStyle'):
            element.set('val', 'marker')
    if kind in ('bar', 'column') and item.get('show_values', True):
        plot = chart.plots[0]
        plot.has_data_labels = True
        plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
        plot.data_labels.font.size = Pt(14)
    # python-pptx templates can emit signed axis ids; normalize both ends of
    # every reference to UInt32 for stricter PPTX importers.
    for element in chart._chartSpace.iter():
        if element.tag in ('{http://schemas.openxmlformats.org/drawingml/2006/chart}axId',
                           '{http://schemas.openxmlformats.org/drawingml/2006/chart}crossAx'):
            identifier = int(element.get('val'))
            if identifier < 0:
                element.set('val', str(identifier + 2 ** 32))
    return shape


def add_table(slide, item, palette):
    rows = item['rows']
    table_shape = slide.shapes.add_table(len(rows), len(rows[0]), *(Inches(value) for value in item['box']))
    table_shape.name = item['id']
    table = table_shape.table
    if item.get('column_weights'):
        weights = item['column_weights']
        if len(weights) != len(rows[0]) or any(value <= 0 for value in weights):
            raise ValueError('Column weights must be positive and match column count')
        for column, weight in zip(table.columns, weights):
            column.width = Inches(item['box'][2] * weight / sum(weights))
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.text = str(value)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(palette[0] if row_index == 0 else ('EDF3F6' if row_index % 2 else 'FFFFFF'))
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(item.get('font_size', 18))
                    run.font.bold = row_index == 0
                    run.font.color.rgb = RGBColor.from_string('FFFFFF' if row_index == 0 else '203446')
    return table_shape


def build(data, output):
    validate(data)
    output = Path(output)
    if output.exists():
        raise FileExistsError('Use a new output path')
    output.parent.mkdir(parents=True, exist_ok=True)
    palette = PALETTES[data.get('profile', 'research')]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for index, spec in enumerate(data['slides'], 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        textbox(slide, 'title', spec['title'], [.6, .35, 12.1, .8], 30, palette[0], True)
        if spec.get('takeaway'):
            textbox(slide, 'takeaway', spec['takeaway'], [.65, 1.16, 12, .65], 20)
        for item in spec['items']:
            (add_chart if item['type'] == 'chart' else add_table)(slide, item, palette)
        sources = ' | '.join(f"{item['id']}: {item['source']}" for item in spec['items'])
        textbox(slide, 'source', sources, [.65, 6.87, 11.6, .35], 11, '566571')
        textbox(slide, 'page', str(index), [12.35, 6.87, .4, .35], 12)
        slide.notes_slide.notes_text_frame.text = spec.get('notes', '') + '\n来源：' + sources
    intermediate = output.with_name(output.stem + '.font-source.pptx')
    if intermediate.exists():
        raise FileExistsError(intermediate)
    prs.save(intermediate)
    normalize(intermediate, output)
    intermediate.unlink()
    return output


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--spec', type=Path, required=True)
    parser.add_argument('--out', type=Path, required=True)
    args = parser.parse_args()
    print(build(json.loads(args.spec.read_text(encoding='utf-8')), args.out))
