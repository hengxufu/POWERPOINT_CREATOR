"""Local template inventory, cross-run replacement and same-deck slide selection."""
import argparse
import json
from pathlib import Path
from pptx import Presentation


def paragraphs(shapes):
    for shape in shapes:
        if shape.shape_type == 6:
            yield from paragraphs(shape.shapes)
        elif shape.has_text_frame:
            yield from shape.text_frame.paragraphs
        elif shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    yield from cell.text_frame.paragraphs


def replace(paragraph, tag, value):
    if not tag:
        raise ValueError('Empty replacement tag')
    count = 0
    while tag in paragraph.text:
        runs = list(paragraph.runs)
        text = ''.join(run.text for run in runs)
        start = text.find(tag)
        if start < 0:
            break
        end, offset = start + len(tag), 0
        for run in runs:
            old = run.text
            stop = offset + len(old)
            if offset < end and stop > start:
                prefix = old[:max(0, start-offset)]
                suffix = old[max(0, end-offset):] if end < stop else ''
                run.text = prefix + (value if offset <= start < stop else '') + suffix
            offset = stop
        count += 1
        # Replacements containing their own tag must not loop forever.
        if tag in value:
            break
    return count


def build(source, spec, output):
    source, output = Path(source), Path(output)
    if output.exists() or source.resolve() == output.resolve():
        raise ValueError('Output must be a new file')
    prs = Presentation(source)
    selection = spec.get('slideSelections')
    if selection is not None:
        numbers = [item['slideNumber'] for item in selection]
        if len(set(numbers)) != len(numbers):
            raise ValueError('Duplicate slide selections are not supported')
        ids = list(prs.slides._sldIdLst)
        if not numbers or any(type(n) is not int or not 1 <= n <= len(ids) for n in numbers):
            raise ValueError('Invalid slide selection')
        for node in ids:
            prs.slides._sldIdLst.remove(node)
        for n in numbers:
            prs.slides._sldIdLst.append(ids[n-1])
        # Keep package relationships: selected slides retain their dependency closure.
    counts = []
    for rule in spec.get('textReplacements', []):
        count = 0
        for number, slide in enumerate(prs.slides, 1):
            if rule.get('slideNumbers') and number not in rule['slideNumbers']:
                continue
            for paragraph in paragraphs(slide.shapes):
                count += replace(paragraph, rule['tag'], rule['value'])
        if not count:
            raise ValueError('Replacement tag absent: ' + rule['tag'])
        counts.append({'tag': rule['tag'], 'count': count})
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    return counts


if __name__ == '__main__':
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('source', type=Path)
    p.add_argument('--spec', type=Path)
    p.add_argument('--out', type=Path, required=True)
    a = p.parse_args()
    if a.spec:
        print(json.dumps(build(a.source, json.loads(a.spec.read_text('utf-8-sig')), a.out)))
    else:
        if a.out.exists() or a.out.resolve() == a.source.resolve():
            p.error('Output must be a new file')
        prs = Presentation(a.source)
        inventory = [{'slideNumber': i, 'layout': s.slide_layout.name,
                      'paragraphs': [v.text for v in paragraphs(s.shapes)]}
                     for i, s in enumerate(prs.slides, 1)]
        a.out.write_text(json.dumps(inventory, ensure_ascii=False, indent=2), 'utf-8')
