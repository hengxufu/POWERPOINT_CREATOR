"""Offline synthetic packages: display order and per-chart workbook failures."""
import tempfile
import unittest
import zipfile
from pathlib import Path
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
from audit_editability import audit


class AuditTests(unittest.TestCase):
    def test_display_order_and_workbook_relationship(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / 'synthetic.pptx'
            prs = Presentation()
            prs.slides.add_slide(prs.slide_layouts[6])
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            data = CategoryChartData()
            data.categories = ['Synthetic A', 'Synthetic B']
            data.add_series('Synthetic values', [1, 2])
            slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,
                                   Inches(1), Inches(1), Inches(5), Inches(3), data)
            order = prs.slides._sldIdLst
            order.insert(0, order[-1])
            prs.save(path)
            result = audit(path)
            self.assertEqual([s['part'] for s in result['slides']],
                             ['ppt/slides/slide2.xml', 'ppt/slides/slide1.xml'])
            chart = result['slides'][0]['charts'][0]
            self.assertEqual(chart['status'], 'embedded_workbook_present')
            with zipfile.ZipFile(path) as source:
                package = {name: source.read(name) for name in source.namelist()}
            # Keep an unrelated workbook: a total count must not hide a broken link.
            package['ppt/embeddings/unrelated.xlsx'] = package[chart['workbook_part']]
            del package[chart['workbook_part']]
            broken = Path(temporary) / 'broken.pptx'
            with zipfile.ZipFile(broken, 'w') as output:
                for name, blob in package.items():
                    output.writestr(name, blob)
            result = audit(broken)
            self.assertEqual(result['embedded_workbooks'], 1)
            self.assertEqual(result['slides'][0]['charts'][0]['status'], 'missing_workbook_part')
            package[chart['workbook_part']] = b'not a workbook'
            with zipfile.ZipFile(broken, 'w') as output:
                for name, blob in package.items():
                    output.writestr(name, blob)
            self.assertEqual(audit(broken)['slides'][0]['charts'][0]['status'],
                             'invalid_workbook_package')


if __name__ == '__main__':
    unittest.main()
