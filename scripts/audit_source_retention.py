"""Compare planned visible text and speaker notes with the actual PPTX."""
import argparse
import hashlib
import json
import re
from pathlib import Path
from pptx import Presentation


def normalized(value):
    return re.sub(r'\s+', '', str(value))


def shape_text(shapes):
    result = []
    for shape in shapes:
        if shape.shape_type == 6:
            result.extend(shape_text(shape.shapes))
        elif shape.has_text_frame:
            result.append(shape.text)
        elif shape.has_table:
            result.extend(cell.text for row in shape.table.rows for cell in row.cells)
    return result


def audit(pptx, plan):
    prs = Presentation(pptx)
    findings, pages = [], []
    planned = plan.get('slides', [])
    if len(planned) != len(prs.slides):
        findings.append({'page': None, 'code': 'page_count_mismatch'})
    for number, slide in enumerate(prs.slides, 1):
        texts = shape_text(slide.shapes)
        visible = '\n'.join(texts)
        # Join runs within a paragraph through shape.text before checking tags.
        tags = re.findall(r'\{\{[^{}\n]+\}\}', visible)
        if tags:
            findings.append({'page': number, 'code': 'unresolved_template_tags', 'values': tags})
        notes = ''
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            notes = frame.text if frame is not None else ''
        spec = planned[number - 1] if number <= len(planned) else {}
        expected_notes = spec.get('speaker_notes', '')
        if expected_notes and normalized(expected_notes) != normalized(notes):
            findings.append({'page': number, 'code': 'speaker_notes_mismatch'})
        required = spec.get('required_visible_text', [])
        if not isinstance(required, list) or not all(isinstance(v, str) for v in required):
            raise ValueError('required_visible_text must be a list of strings')
        for value in required:
            if normalized(value) not in normalized(visible):
                findings.append({'page': number, 'code': 'required_text_missing', 'value': value})
        pages.append({'page': number, 'slide_id': spec.get('slide_id'), 'text': visible, 'speaker_notes': notes})
    return {'ready': not findings, 'pptx_sha256': hashlib.sha256(Path(pptx).read_bytes()).hexdigest(),
            'findings': findings, 'slides': pages,
            'limits': ['Text retention does not prove visual readability or factual accuracy.',
                       'Charts, formulas, inherited text and images require separate checks.']}


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--pptx', required=True, type=Path)
    parser.add_argument('--plan', required=True, type=Path)
    parser.add_argument('--out', required=True, type=Path)
    args = parser.parse_args()
    if args.out.resolve() in (args.pptx.resolve(), args.plan.resolve()):
        parser.error('Report must not overwrite input')
    result = audit(args.pptx, json.loads(args.plan.read_text(encoding='utf-8-sig')))
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    raise SystemExit(0 if result['ready'] else 1)
