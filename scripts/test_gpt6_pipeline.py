#!/usr/bin/env python3
"""Offline self-test for PPT session initialization and element composition."""

from __future__ import annotations

import hashlib
import copy
import json
import subprocess
import sys
import tempfile
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation


HERE = Path(__file__).resolve().parent


def dump(path: Path, value: dict) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def run(*args: object, expect_success: bool = True) -> subprocess.CompletedProcess[str]:
    result = subprocess.run(
        [str(arg) for arg in args], text=True, capture_output=True, encoding="utf-8", errors="replace"
    )
    if expect_success and result.returncode != 0:
        raise AssertionError(result.stdout + result.stderr)
    if not expect_success and result.returncode == 0:
        raise AssertionError("Command unexpectedly succeeded")
    return result


def main() -> int:
    with tempfile.TemporaryDirectory(prefix="research-ppt-skill-", dir=HERE.parent) as temporary:
        root = Path(temporary)
        workspace = root / "workspace"
        run(
            sys.executable,
            HERE / "init_ppt_session.py",
            "--out", workspace,
            "--title", "测试汇报",
            "--project-root", root,
        )
        session = json.loads((workspace / "session.json").read_text(encoding="utf-8"))
        assert session["planning_owner"] == "chatgpt-web:gpt-6-astra"
        assert (workspace / "planning" / "gpt6_task_packet.md").is_file()

        source = root / "source.pptx"
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[6])
        prs.save(source)

        asset = root / "asset.png"
        image = Image.new("RGBA", (600, 600), (0, 0, 0, 0))
        ImageDraw.Draw(image).ellipse((60, 60, 540, 540), fill=(16, 114, 108, 255))
        image.save(asset)
        digest = hashlib.sha256(asset.read_bytes()).hexdigest()

        job = root / "job"
        job.mkdir()
        dump(job / "job.json", {
            "status": "accepted",
            "spec": {"kind": "icon", "slide_id": "slide-01", "object_id": "sensor"},
            "asset_file": str(asset),
            "sha256": digest,
            "technical_checks": {"width": 600, "height": 600, "subject_bbox_px": [60, 60, 540, 540]},
        })
        run(sys.executable, HERE / "asset_job.py", "placement", "--job", job,
            "--box", 1, 1, 2, 2, "--canvas", 10, 7.5)
        manifest = root / "manifest.json"
        dump(manifest, {
            "schema": "ppt_element_manifest_v1",
            "max_generated_asset_area_fraction": 0.55,
            "elements": [{"kind": "icon", "slide_number": 1, "job_dir": str(job)}],
        })
        output = root / "output.pptx"
        report = root / "report.json"
        run(sys.executable, HERE / "compose_elements.py", "--pptx", source, "--manifest", manifest, "--out", output, "--report", report)
        result = Presentation(output)
        pictures = [shape for shape in result.slides[0].shapes if shape.name == "slide-01/sensor"]
        assert len(pictures) == 1
        assert hashlib.sha256(pictures[0].image.blob).hexdigest() == digest
        assert json.loads(report.read_text(encoding="utf-8"))["whole_slide_generation_used"] is False
        assert json.loads(report.read_text(encoding="utf-8"))["manifest_sha256"] == hashlib.sha256(manifest.read_bytes()).hexdigest()

        valid_placement = json.loads((job / "placement.json").read_text(encoding="utf-8"))
        mutations = [
            ("wrong-job-hash", "asset_sha256", "0" * 64, "placement hash differs"),
            ("wrong-object", "object_id", "other", "placement object_id mismatch"),
            ("stretched", "image_box_inches", [1, 1, 2, 1], "would stretch"),
            ("nonfinite", "image_box_inches", [float('nan'), 1, 2, 2], "non-finite"),
            ("oversized", "image_box_inches", [0, 0, 7, 7], "area limit"),
        ]
        for name, field, value, message in mutations:
            broken = copy.deepcopy(valid_placement)
            broken[field] = value
            dump(job / "placement.json", broken)
            rejected = run(sys.executable, HERE / "compose_elements.py", "--pptx", source,
                           "--manifest", manifest, "--out", root / f"{name}.pptx",
                           "--report", root / f"{name}.json", expect_success=False)
            assert message in rejected.stderr, rejected.stderr
            assert not (root / f"{name}.pptx").exists()
        dump(job / "placement.json", valid_placement)

        replacement = root / "replacement.png"
        Image.new("RGBA", (600, 600), (255, 0, 0, 255)).save(replacement)
        broken = {**valid_placement, "asset_file": str(replacement),
                  "asset_sha256": hashlib.sha256(replacement.read_bytes()).hexdigest()}
        dump(job / "placement.json", broken)
        rejected = run(sys.executable, HERE / "compose_elements.py", "--pptx", source,
                       "--manifest", manifest, "--out", root / "substitution.pptx",
                       "--report", root / "substitution.json", expect_success=False)
        assert "differs from accepted job" in rejected.stderr
        dump(job / "placement.json", valid_placement)

        rejected = run(sys.executable, HERE / "compose_elements.py", "--pptx", source,
                       "--manifest", manifest, "--out", root / "same.pptx",
                       "--report", root / "same.pptx", expect_success=False)
        assert "different paths" in rejected.stderr
        assert not (root / "same.pptx").exists()

        # A matching name on another page would make later integration ambiguous.
        duplicate_source = root / "duplicate-source.pptx"
        duplicate = Presentation(str(source))
        second = duplicate.slides.add_slide(duplicate.slide_layouts[6])
        second.shapes.add_picture(str(asset), 0, 0).name = "slide-01/sensor"
        duplicate.save(duplicate_source)
        rejected = run(sys.executable, HERE / "compose_elements.py", "--pptx", duplicate_source,
                       "--manifest", manifest, "--out", root / "duplicate.pptx",
                       "--report", root / "duplicate.json", expect_success=False)
        assert "already exists in deck" in rejected.stderr

        bad_manifest = root / "bad-manifest.json"
        dump(bad_manifest, {
            "schema": "ppt_element_manifest_v1",
            "elements": [{"kind": "full_slide", "slide_number": 1, "job_dir": str(job)}],
        })
        run(
            sys.executable, HERE / "compose_elements.py",
            "--pptx", source,
            "--manifest", bad_manifest,
            "--out", root / "forbidden.pptx",
            "--report", root / "forbidden.json",
            expect_success=False,
        )
    print("offline pipeline self-test passed")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
