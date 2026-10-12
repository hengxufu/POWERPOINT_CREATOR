"""Offline synthetic decks testing real local-template identity assets, not web evidence."""
import copy
import json
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

from beihang_theme import PACK, PREFIX, apply, audit, contract, sha
from audit_element_density import audit as density
from auto_workflow import prepare, load, infer


class ThemeTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name).resolve()
        self.input = self.root / 'input.pptx'
        self.output = self.root / 'themed.pptx'
        self.plan = {'design_system': {'theme_profile': 'beihang-blue'}, 'slides': [
            {'slide_number': 1, 'section': '研究背景'}, {'slide_number': 2, 'section': '方法与验证'}]}
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        for _ in range(2):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            text = slide.shapes.add_textbox(Inches(1), Inches(1.2), Inches(7), Inches(.8))
            text.text = 'Synthetic scientific content; offline fixture only'
        prs.save(self.input)
        self.original = sha(self.input)

    def tearDown(self):
        self.temp.cleanup()

    def test_pack_preserves_original_bytes_and_provenance(self):
        manifest = load(PACK / 'manifest.json')
        self.assertEqual(len(manifest['templates']), 12)
        for record in manifest['assets'].values():
            self.assertEqual(record['sha256'], sha(PACK / record['file']))
            self.assertEqual(record['transform'], 'none')

    def test_every_page_has_four_objects_and_preserves_content_and_input(self):
        result = apply(self.input, self.plan, self.output)
        self.assertTrue(result['ready'], result['problems'])
        self.assertEqual(sha(self.input), self.original)
        prs = Presentation(self.output)
        for slide in prs.slides:
            self.assertEqual(sum(shape.name.startswith(PREFIX) for shape in slide.shapes), 4)
            self.assertIn('Synthetic scientific content', slide.shapes[0].text)
        self.assertEqual([p['element_count'] for p in density(self.output)['slides']], [5, 5])

    def test_reapply_does_not_accumulate_theme_objects(self):
        apply(self.input, self.plan, self.output)
        second = self.root / 'second.pptx'
        result = apply(self.output, self.plan, second)
        self.assertTrue(result['ready'])
        self.assertEqual(len(Presentation(second).slides[0].shapes), 5)

    def mutate(self, callback):
        apply(self.input, self.plan, self.output)
        prs = Presentation(self.output)
        callback(prs)
        prs.save(self.output)
        return audit(self.output, self.plan)

    def test_missing_identity_fails(self):
        def change(prs):
            shape = next(s for s in prs.slides[1].shapes if s.name == PREFIX + 'identity')
            prs.slides[1].shapes._spTree.remove(shape._element)
        self.assertFalse(self.mutate(change)['ready'])

    def test_renamed_picture_cannot_fake_source_identity(self):
        def change(prs):
            slide = prs.slides[0]
            shape = next(s for s in slide.shapes if s.name == PREFIX + 'identity')
            box = [shape.left, shape.top, shape.width, shape.height]
            slide.shapes._spTree.remove(shape._element)
            slide.shapes.add_picture(str(PACK / 'campus-panorama.png'), *box).name = PREFIX + 'identity'
        self.assertTrue(any('not original' in p for p in self.mutate(change)['problems']))

    def test_hidden_cropped_distorted_and_wrong_text_fail(self):
        callbacks = [
            lambda prs: next(s for s in prs.slides[0].shapes if s.name == PREFIX+'identity')._element.xpath('.//p:cNvPr')[0].set('hidden', '1'),
            lambda prs: setattr(next(s for s in prs.slides[0].shapes if s.name == PREFIX+'identity'), 'crop_left', .1),
            lambda prs: setattr(next(s for s in prs.slides[0].shapes if s.name == PREFIX+'identity'), 'height', Inches(.7)),
            lambda prs: setattr(next(s for s in prs.slides[0].shapes if s.name == PREFIX+'page'), 'text', '99'),
        ]
        for callback in callbacks:
            with self.subTest(callback=callback):
                if self.output.exists():
                    self.output.unlink()
                self.assertFalse(self.mutate(callback)['ready'])

    def test_foreground_collision_is_not_silently_passed(self):
        prs = Presentation(self.input)
        prs.slides[0].shapes.add_textbox(Inches(10.7), Inches(.2), Inches(2), Inches(.5)).text = 'Conflicting content'
        prs.save(self.input)
        result = apply(self.input, self.plan, self.output)
        self.assertFalse(result['ready'])
        self.assertTrue(result['collisions'])

    def test_small_canvas_and_alternative_theme(self):
        prs = Presentation(self.input)
        prs.slide_width, prs.slide_height = Inches(10), Inches(7.5)
        prs.save(self.input)
        plan = copy.deepcopy(self.plan)
        plan['design_system']['theme_profile'] = 'beihang-green'
        self.assertTrue(apply(self.input, plan, self.output, 'beihang-green')['ready'])
        self.assertFalse(audit(self.output, plan, 'beihang-blue')['ready'])

    def test_invalid_plan_or_same_output_rejected(self):
        bad = copy.deepcopy(self.plan)
        bad['slides'][0].pop('section')
        with self.assertRaises(ValueError):
            apply(self.input, bad, self.output)
        self.assertFalse(self.output.exists())
        with self.assertRaises(ValueError):
            apply(self.input, self.plan, self.input)
        self.assertEqual(sha(self.input), self.original)

    def test_default_theme_automatic_resume_and_explicit_opt_out(self):
        # No web transport or model-generated fixtures are used by prepare.
        result = prepare(self.root)
        workspace = Path(result['workspace'])
        self.assertEqual(load(workspace / 'session.json')['theme_profile'], 'beihang-blue')
        self.assertEqual(load(workspace / 'planning/theme-contract.json')['recurring_elements'], ['identity', 'university', 'section', 'page'])
        self.assertIn('wordmark-blue.png', Path(result['brief']).read_text(encoding='utf-8'))
        prepare(self.root, theme_profile='beihang-ginkgo')
        prepare(self.root)
        self.assertEqual(load(workspace / 'session.json')['theme_profile'], 'beihang-ginkgo')
        with self.assertRaises(ValueError):
            infer(self.root, workspace, theme_profile='none')
        prepare(self.root, theme_profile='none', theme_reason='Synthetic explicit user instruction: use another school')
        self.assertEqual(load(workspace / 'composition_contract.json')['theme_profile'], 'none')


if __name__ == '__main__':
    unittest.main()
