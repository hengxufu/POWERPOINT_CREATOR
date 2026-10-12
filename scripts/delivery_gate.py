"""Reconcile the final PPTX with its plan, web jobs and page review evidence.

This is an evidence gate, not an aesthetic scoring model or browser driver.
Report pass means recorded evidence is complete, not independently proven truthful.
"""
import argparse
import hashlib
import json
from pathlib import Path
from pptx import Presentation


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def load(path):
    return json.loads(Path(path).read_text(encoding='utf-8-sig'))


def resolve(base, value):
    value = Path(value)
    return value if value.is_absolute() else base / value


def check(workspace, pptx):
    workspace, pptx = Path(workspace), Path(pptx)
    problems, warnings = [], []

    def require(condition, message):
        if not condition:
            problems.append(message)

    try:
        plan = load(workspace / 'planning/deck_plan.json')
        review = load(workspace / 'qa/delivery_review.json')
        manifest = load(workspace / 'assets/asset_manifest.json')
        prs = Presentation(pptx)
    except (OSError, ValueError, KeyError) as error:
        return {'ready': False, 'status': 'draft', 'problems': [str(error)], 'warnings': []}
    deck_hash = digest(pptx)
    theme_declarations = []
    for relative in ('session.json', 'composition_contract.json', 'planning/theme-contract.json', 'planning/auto-brief.json'):
        path = workspace / relative
        if path.is_file():
            declaration = load(path)
            if declaration.get('theme_profile'):
                theme_declarations.append(declaration['theme_profile'])
    if plan.get('design_system', {}).get('theme_profile'):
        theme_declarations.append(plan['design_system']['theme_profile'])
    theme_audit = None
    if theme_declarations:
        require(len(set(theme_declarations)) == 1, 'Conflicting theme profiles in planning/session/composition contracts')
        profile = theme_declarations[0]
        if profile.startswith('beihang-'):
            from beihang_theme import audit as audit_theme
            try:
                theme_audit = audit_theme(pptx, plan, profile)
                for problem in theme_audit['problems']:
                    require(False, 'Theme: ' + problem)
            except (ValueError, KeyError, OSError) as error:
                require(False, 'Theme: ' + str(error))
        elif profile == 'none':
            session = load(workspace / 'session.json') if (workspace / 'session.json').is_file() else {}
            require(bool(session.get('theme_reason', '').strip()), 'Theme disabled without a specific recorded user instruction')
        else:
            require(False, 'Unsupported theme profile')
    from audit_element_density import audit as audit_density
    density = audit_density(pptx, plan=plan)
    for page in density['slides']:
        require(page['deficit'] == 0, f"Page {page['slide_number']}: {page['element_count']}/40 elements; generate and compose {page['deficit']} independent web elements")
    require(review.get('pptx_sha256') == deck_hash, 'Review belongs to another PPTX revision')
    slides = plan.get('slides', [])
    require(len(slides) == len(prs.slides) and bool(slides), 'Plan does not cover every actual slide')
    numbers = [slide.get('slide_number') for slide in slides]
    require(numbers == list(range(1, len(prs.slides) + 1)), 'Plan pages must follow actual display order')
    require(review.get('unresolved_issues') == [], 'Unresolved issues have not been cleared')
    planning = review.get('web_planning', {})
    offline = review.get('offline_user_authorization', {})
    if offline.get('instruction') and offline.get('reason'):
        warnings.append('User-authorized offline mode; do not claim web generation/planning')
    else:
        from bridge_session import check as check_connection
        connection = check_connection(workspace)
        for problem in connection['problems']:
            require(False, 'Connection: ' + problem)
        connection_record = load(workspace / 'bridge-connect.json') if (workspace / 'bridge-connect.json').is_file() else {}
        if connection_record.get('unverified_planning_user_authorization'):
            warnings.append('Planning used user-authorized web chat with unverified exact model; do not label it verified GPT-6.')
        if connection_record.get('unverified_image_user_authorization'):
            warnings.append('Image generation used user-authorized web capability with unverified exact engine; do not label it verified Image 2.5.')
        response = planning.get('response_file')
        require(bool(response), 'Missing actual web planning response')
        if response:
            path = resolve(workspace, response)
            require(path.is_file() and bool(path.read_text(encoding='utf-8').strip()), 'Web planning response missing or empty')
        require(str(planning.get('session_url', '')).startswith('https://chatgpt.com/c/'), 'Missing observed web planning conversation')
    reviewed_pages = review.get('slides', [])
    require([page.get('slide_number') for page in reviewed_pages] == numbers, 'Missing, duplicated or out-of-order page reviews')
    page_reviews = {page.get('slide_number'): page for page in reviewed_pages}
    requests, jobs = {}, {}
    for element in manifest.get('elements', []):
        key = element.get('request_id')
        require(bool(key) and key not in jobs, 'Manifest must have unique request_id values')
        if key:
            jobs[key] = element
    layouts = []
    for slide in slides:
        number = slide.get('slide_number')
        label = f'Page {number}'
        for field in ('title', 'takeaway', 'layout_family', 'visual_rationale'):
            require(bool(str(slide.get(field, '')).strip()), f'{label}: missing {field}')
        layouts.append(slide.get('layout_family'))
        blocks = slide.get('content_blocks', [])
        require(bool(blocks) and all(block.get('text') or block.get('data_ref') or block.get('source_ref') for block in blocks), f'{label}: content is empty or only headings')
        record = page_reviews.get(number, {})
        for field in ('content_review', 'visual_review', 'evidence_review'):
            item = record.get(field, {})
            require(item.get('status') == 'pass' and bool(item.get('note', '').strip()), f'{label}: missing actual {field}')
        require(record.get('rendered_from_sha256') == deck_hash, f'{label}: render is stale or not tied to final deck')
        render = record.get('render_file')
        if render:
            path = resolve(workspace, render)
            require(path.is_file() and digest(path) == record.get('render_sha256'), f'{label}: render file/hash mismatch')
        else:
            require(False, f'{label}: no page render')
        for request in slide.get('asset_requests', []):
            key = request.get('request_id')
            require(bool(key) and key not in requests, f'{label}: asset request_id missing or duplicate')
            if not key:
                continue
            requests[key] = number
            if request.get('decision') == 'native_or_source':
                require(bool(request.get('decision_reason')) and bool(request.get('replacement_object')), f'{label}/{key}: unjustified image omission')
                if type(number) is int and 1 <= number <= len(prs.slides):
                    require(any(shape.name == request.get('replacement_object') for shape in prs.slides[number - 1].shapes), f'{label}/{key}: replacement object absent')
                continue
            element = jobs.get(key)
            require(element is not None, f'{label}/{key}: required web element was not composed')
            if element is None:
                continue
            require(element.get('slide_number') == number, f'{label}/{key}: wrong destination page')
            try:
                job_dir = resolve(workspace / 'assets', element['job_dir'])
                job = load(job_dir / 'job.json')
                require(job.get('transport') == 'chatgpt_web', f'{key}: requested web route was substituted')
                require(job.get('status') in ('accepted', 'integrated'), f'{key}: unfinished job')
                require(any(event.get('status') == 'submitted' for event in job.get('history', [])), f'{key}: no recorded web submission')
                require(bool(job.get('visual_review', {}).get('note')), f'{key}: asset has not been visually reviewed')
                if not job.get('verified_engine'):
                    warnings.append(f'{key}: actual image engine unverified; label it honestly')
                name = job['spec']['slide_id'] + '/' + job['spec']['object_id']
                shapes = [shape for shape in prs.slides[number - 1].shapes if shape.name == name] if type(number) is int and 1 <= number <= len(prs.slides) else []
                require(len(shapes) == 1 and hasattr(shapes[0], 'image') and hashlib.sha256(shapes[0].image.blob).hexdigest() == job.get('sha256'), f'{key}: accepted image not found on destination page')
            except (OSError, ValueError, KeyError) as error:
                require(False, f'{key}: {error}')
    require(set(jobs) <= set(requests), 'Manifest includes elements not considered in the plan')
    if len(slides) >= 5 and len(set(layouts)) < 3:
        require(bool(review.get('layout_repetition_reason')), 'Repetitive layout needs revision or a content-based explanation')
    require(bool(requests) or bool(review.get('no_generated_assets_reason')), 'No image requests: explain the visual strategy instead of silently skipping images')
    return {'ready': not problems, 'status': 'review_evidence_complete' if not problems else 'draft',
            'pptx_sha256': deck_hash, 'problems': problems, 'warnings': warnings,
            'element_density': density,
            'beihang_theme': theme_audit,
            'limits': ['Review notes are human/agent attestations, not automatic proof of aesthetics or factual accuracy.',
                       'Run editability, typography and source audits separately.']}


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--workspace', required=True, type=Path)
    parser.add_argument('--pptx', required=True, type=Path)
    parser.add_argument('--out', required=True, type=Path)
    args = parser.parse_args()
    if args.out.resolve() == args.pptx.resolve():
        parser.error('Report must not overwrite PPTX')
    result = check(args.workspace, args.pptx)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(result, ensure_ascii=False))
    raise SystemExit(0 if result['ready'] else 1)
