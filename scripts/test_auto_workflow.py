"""Offline automation regressions; all sources/plans/jobs here are synthetic."""
import json
import tempfile
import unittest
from pathlib import Path

from auto_workflow import choose_workspace, infer, inventory, prepare, status, write


class AutomationTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name).resolve()
        self.workspace = self.root / 'outputs/research-ppt'

    def tearDown(self):
        self.temporary.cleanup()

    def test_minimal_request_initializes_draft_with_packet(self):
        result = prepare(self.root)
        self.assertTrue(Path(result['planning_packet']).is_file())
        brief = json.loads((self.workspace / 'planning/auto-brief.json').read_text(encoding='utf-8'))
        self.assertEqual(brief['duration_min'], 20)
        self.assertEqual(brief['profile'], 'research')
        self.assertTrue(brief['evidence_gaps'])
        self.assertEqual(status(self.workspace)['next_action'], 'observe_browser_and_check_connection')
        self.assertFalse((self.workspace / 'bridge-connect.json').exists())

    def test_explicit_values_and_resume_preserve_original_request(self):
        prepare(self.root, request='做一份《合成题目》的国奖答辩，8分钟', duration=12)
        resumed = prepare(self.root)
        brief = json.loads((self.workspace / 'planning/auto-brief.json').read_text(encoding='utf-8'))
        self.assertEqual(brief['duration_min'], 12)
        self.assertEqual(brief['title'], '合成题目')
        self.assertEqual(brief['profile'], 'scholarship')
        self.assertIn('8分钟', Path(resumed['brief']).read_text(encoding='utf-8'))
        write(self.workspace / 'asset-jobs/sentinel.json', {'status': 'submission_unknown'})
        again = prepare(self.root)
        self.assertEqual(again['planning_packet'], resumed['planning_packet'])
        self.assertTrue((self.workspace / 'asset-jobs/sentinel.json').exists())

    def test_metadata_inventory_is_bounded_and_excludes_credentials(self):
        (self.root / '资料').mkdir()
        (self.root / '资料/合成报告.pdf').write_bytes(b'Not an actual PDF; metadata fixture only')
        (self.root / 'password.txt').write_text('SYNTHETIC', encoding='utf-8')
        (self.root / 'AGENTS.md').write_text('SYNTHETIC', encoding='utf-8')
        entries = inventory(self.root)
        self.assertEqual(len(entries), 1)
        self.assertFalse(entries[0]['content_read'])
        brief = infer(self.root, self.workspace)
        self.assertEqual(brief['title'], '合成报告')
        self.assertEqual(brief['provenance']['title'], 'material_filename')

    def test_multiple_workspaces_require_selection(self):
        for name in ('first', 'second'):
            write(self.root / 'outputs' / name / 'session.json', {})
        with self.assertRaisesRegex(ValueError, 'Multiple PPT workspaces'):
            choose_workspace(self.root)
        self.assertEqual(choose_workspace(self.root, 'outputs/first'), self.root / 'outputs/first')

    def test_workspace_cannot_escape_project(self):
        with self.assertRaisesRegex(ValueError, 'child directory'):
            prepare(self.root, self.root.parent / 'escaped-ppt')
        self.assertFalse(self.workspace.exists())

    def test_unknown_submission_is_never_queued_for_resubmission(self):
        prepare(self.root)
        plan = {'slides': [{'slide_number': 1, 'slide_id': 's1', 'title': 'Synthetic',
                           'takeaway': 'Synthetic', 'layout_family': 'figure-analysis',
                           'visual_rationale': 'Test only',
                           'content_blocks': [{'text': 'Synthetic only', 'box_inches': [1, 1, 3, 2]}],
                           'asset_requests': [{'request_id': 'r1', 'kind': 'icon', 'object_id': 'icon',
                                               'single_subject_prompt': 'Synthetic icon'}]}]}
        write(self.workspace / 'planning/deck_plan.json', plan)
        write(self.workspace / 'asset-jobs/v1/queue.json', {'jobs': [{'request_id': 'r1', 'job_dir': 'r1'}]})
        write(self.workspace / 'asset-jobs/v1/r1/job.json', {'status': 'submission_unknown'})
        result = status(self.workspace)
        self.assertEqual(result['asset_actions'][0]['action'], 'inspect_original_conversation_do_not_resubmit')
        self.assertEqual(result['delivery_status'], 'draft')
        (self.workspace / 'asset-jobs/v1/queue.json').unlink()
        recovered = status(self.workspace)
        self.assertEqual(recovered['asset_actions'][0]['state'], 'submission_unknown')
        self.assertEqual(recovered['asset_actions'][0]['action'], 'inspect_original_conversation_do_not_resubmit')
        from pptx import Presentation
        deck = Presentation()
        deck.slides.add_slide(deck.slide_layouts[6])
        deck.save(self.workspace / 'build/draft.pptx')
        counted = status(self.workspace)
        self.assertEqual(counted['element_density']['total_deficit'], 40)
        self.assertEqual(counted['next_action'], 'observe_browser_and_check_connection')

    def test_explicit_profile_change_updates_scenario(self):
        prepare(self.root)
        prepare(self.root, profile='challenge')
        brief = json.loads((self.workspace / 'planning/auto-brief.json').read_text(encoding='utf-8'))
        self.assertEqual(brief['profile'], 'challenge')
        self.assertEqual(brief['scenario'], '挑战杯答辩')

    def test_invalid_duration_does_not_create_workspace(self):
        with self.assertRaisesRegex(ValueError, 'Duration'):
            prepare(self.root, duration=0)
        self.assertFalse(self.workspace.exists())

    def test_unknown_planning_submission_is_preserved_before_new_plan(self):
        prepare(self.root)
        write(self.workspace / 'planning/web-design-dispatch.json',
              {'state': 'submission_unknown', 'send_action_attempted': True,
               'actual_planning_response_received': False})
        resumed = status(self.workspace)
        self.assertEqual(resumed['next_action'], 'recover_original_planning_conversation_do_not_resubmit')
        self.assertFalse(resumed['planning_dispatch']['actual_planning_response_received'])

    def test_material_content_reaches_both_packets_and_controls_resume(self):
        first = prepare(self.root)
        self.assertFalse(first['materials_prepared'])
        material = self.workspace/'sources/context.md'
        material.write_text('SYNTHETIC project context and blue icon appearance', encoding='utf-8')
        selection = Path(first['material_selection'])
        write(selection, {'schema': 'web_material_selection_v1', 'coverage': {'gpt': 'Synthetic brief', 'image': 'Synthetic style'},
                          'missing': [], 'materials': [{'id': 'context', 'path': str(material.relative_to(self.root)),
                                                       'targets': ['gpt', 'image'], 'mode': 'inline', 'purpose': 'Fixture', 'locator': 'Whole fixture'}]})
        prepared = prepare(self.root)
        self.assertTrue(prepared['materials_prepared'])
        from material_handoff import check as material_check, record, digest
        for key in ('planning_material_bundle', 'image_material_bundle'):
            bundle = Path(prepared[key])
            self.assertIn('SYNTHETIC project context', (bundle/'message.txt').read_text(encoding='utf-8'))
            self.assertTrue(material_check(bundle, require_transfer=False)['ready'])
            self.assertFalse(material_check(bundle)['ready'])
        from bridge_session import prepare as connection_prepare
        connection = connection_prepare(self.workspace, self.root)
        connection.update(local_profile_configured=True, browser_connected=True, image_tool_visible=True,
                          planning_model='GPT-6', image_model='Image-2.5', browser_route='builtin_browser',
                          route_reason='SYNTHETIC ONLY', observed_connection_id=connection['connection_id'],
                          evidence=[{'file': 'sources/context.md', 'sha256': digest(material)}])
        write(self.workspace/'bridge-connect.json', connection)
        self.assertEqual(status(self.workspace)['next_action'], 'transfer_planning_materials')
        observation = self.workspace/'qa/synthetic-transfer.json'
        write(observation, {'route': 'builtin_browser', 'session_url': 'https://chatgpt.com/c/SYNTHETIC', 'message_visible': True,
                            'note': 'SYNTHETIC ONLY', 'evidence_file': str(material.relative_to(self.root)),
                            'materials': [{'id': 'context', 'mode': 'inline', 'visible_note': 'Synthetic visible source'}]})
        record(Path(prepared['planning_material_bundle']), observation)
        self.assertEqual(status(self.workspace)['next_action'], 'receive_actual_web_plan')
        material.write_text('CHANGED SYNTHETIC MATERIAL', encoding='utf-8')
        connection['evidence'][0]['sha256'] = digest(material)
        write(self.workspace/'bridge-connect.json', connection)
        self.assertEqual(status(self.workspace)['next_action'], 'prepare_planning_materials')
        updated = prepare(self.root)
        self.assertNotEqual(updated['planning_material_bundle'], prepared['planning_material_bundle'])
        self.assertTrue((Path(prepared['planning_material_bundle'])/'transfer.json').exists())


if __name__ == '__main__':
    unittest.main()
