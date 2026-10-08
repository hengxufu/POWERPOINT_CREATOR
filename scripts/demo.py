"""Small source-first native PPTX demo; not a replacement for existing deck engines."""
import argparse
import json
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.chart.data import CategoryChartData
from pptx.oxml.xmlchemy import OxmlElement

FONT = 'Microsoft YaHei'
INK, MUTED, TEAL, LIGHT = '162B3A', '556575', '087F8C', 'E9F3F4'

def source():
    return {'title': '中国科研汇报｜可编辑能力样例', 'slides': [
        {'id': 's01', 'role': 'cover', 'objects': [
            {'id': 'title', 'type': 'text', 'text': '让证据成为\n科研汇报的主角', 'box': [.7, 1.45, 10.8, 1.7], 'size': 44},
            {'id': 'subtitle', 'type': 'text', 'text': '中文科研汇报 · 可编辑对象 · 持续交互打磨', 'box': [.75, 3.6, 11, .55], 'size': 22},
            {'id': 'disclosure', 'type': 'text', 'text': '制作能力演示｜所有数值为示例，不代表任何研究结果', 'box': [.75, 5.45, 11, .55], 'size': 16}],
         'notes': '演示设计和编辑能力；内容不是科学实验结论。'},
        {'id': 's02', 'role': 'question', 'title': '先确定专家需要判断的问题', 'objects': [
            {'id': 'question', 'type': 'text', 'text': '这项工作解决了什么问题？', 'box': [.7, 1.6, 11.8, .8], 'size': 34},
            {'id': 'evidence', 'type': 'text', 'text': '证据：研究设计、关键结果与可核查来源', 'box': [.75, 3.0, 11.7, .65], 'size': 24},
            {'id': 'boundary', 'type': 'text', 'text': '边界：已完成的发现、推断与未来假设分别表达', 'box': [.75, 4.05, 11.7, .65], 'size': 24}],
         'notes': '这是建议的内容组织方式；实际汇报按组会、答辩或项目评审要求调整。'},
        {'id': 's03', 'role': 'result', 'title': '结果页先让图表可读，再解释结论', 'objects': [
            {'id': 'main-chart', 'type': 'chart', 'box': [.65, 1.55, 8.0, 4.75], 'categories': ['条件 A', '条件 B', '条件 C'], 'values': [42, 57, 68]},
            {'id': 'interpretation', 'type': 'text', 'text': '仅展示图表编辑能力\n\n可修改数值、标签和配色\n\n示例未提供统计检验，\n不作显著性判断', 'box': [9.05, 1.95, 3.35, 3.7], 'size': 21},
            {'id': 'source-label', 'type': 'text', 'text': '数据：本地合成示例｜单位：演示得分｜无实测样本', 'box': [.75, 6.35, 11.6, .45], 'size': 14}],
         'notes': '数据为合成示例。图表具有内嵌 Excel 数据；不推断实验结论。'},
        {'id': 's04', 'role': 'table', 'title': '把审稿式追问转化为可核查的证据表', 'objects': [
            {'id': 'evidence-table', 'type': 'table', 'box': [.7, 1.75, 11.85, 3.8], 'rows': [
                ['专家的问题', '需要的证据', '汇报处理'],
                ['结果可靠吗？', '对照、重复与统计口径', '注明条件与局限'],
                ['创新在哪里？', '与现有研究的同口径比较', '界定增量贡献'],
                ['下一步能完成吗？', '研究基础与替代方案', '展示验证路径']]},
            {'id': 'note', 'type': 'text', 'text': '文字与单元格均为原生对象；可在 PowerPoint 中修改', 'box': [.75, 6.0, 11.6, .6], 'size': 20}],
         'notes': '表格是表达建议示例，不代表具体项目评审规则。'},
        {'id': 's05', 'role': 'close', 'title': '精细修改要落在稳定对象上', 'objects': [
            {'id': 'instruction', 'type': 'text', 'text': '“把第 3 页右侧解释改短一些”', 'box': [.75, 1.75, 11.65, 1.0], 'size': 34},
            {'id': 'mapping', 'type': 'text', 'text': '定位 s03 / interpretation → 修改源 → 重建 → 检查', 'box': [.75, 3.3, 11.6, .7], 'size': 23},
            {'id': 'limits', 'type': 'text', 'text': '结构检查不能替代渲染、事实核对与真实放映检查', 'box': [.75, 4.7, 11.6, .7], 'size': 22}],
         'notes': '保留确认过的内容与版本；本样例以源文件重建，不实现自然语言解析器。'}]}

def font(run, size, color=INK, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._r.get_or_add_rPr()
    for tag in ('a:ea', 'a:cs'):
        element = OxmlElement(tag)
        element.set('typeface', FONT)
        rpr.append(element)

def text(slide, name, value, box, size, color=INK, bold=False):
    shape = slide.shapes.add_textbox(*(Inches(n) for n in box))
    shape.name = name
    shape.text_frame.word_wrap = True
    shape.text_frame.margin_left = shape.text_frame.margin_right = Inches(0)
    shape.text_frame.margin_top = shape.text_frame.margin_bottom = Inches(0)
    for i, line in enumerate(value.split('\n')):
        p = shape.text_frame.paragraphs[0] if i == 0 else shape.text_frame.add_paragraph()
        p.space_after = Pt(8)
        font(p.add_run(), size, color, bold)
        p.runs[-1].text = line
    return shape

def build(data, output):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for idx, spec in enumerate(data['slides'], 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        if spec['role'] == 'cover':
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor.from_string(INK)
        if spec.get('title'):
            text(slide, spec['id']+'/title', spec['title'], [.7, .5, 11.9, .85], 30, bold=True)
        for obj in spec['objects']:
            name = spec['id']+'/'+obj['id']
            if obj['type'] == 'text':
                text(slide, name, obj['text'], obj['box'], obj['size'], 'FFFFFF' if spec['role']=='cover' else INK)
            elif obj['type'] == 'chart':
                chartdata = CategoryChartData()
                chartdata.categories = obj['categories']
                chartdata.add_series('演示得分（示例）', obj['values'])
                shape = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, *(Inches(n) for n in obj['box']), chartdata)
                shape.name = name
                chart = shape.chart
                chart.has_legend = False
                chart.category_axis.tick_labels.font.name = FONT
                chart.category_axis.tick_labels.font.size = Pt(18)
                chart.value_axis.tick_labels.font.size = Pt(14)
                chart.value_axis.minimum_scale, chart.value_axis.maximum_scale = 0, 100
                plot = chart.plots[0]
                plot.has_data_labels = True
                plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
                plot.data_labels.font.size = Pt(18)
                chart.series[0].format.fill.solid()
                chart.series[0].format.fill.fore_color.rgb = RGBColor.from_string(TEAL)
            elif obj['type'] == 'table':
                rows = obj['rows']
                shape = slide.shapes.add_table(len(rows), len(rows[0]), *(Inches(n) for n in obj['box']))
                shape.name = name
                for i, row in enumerate(rows):
                    for j, value in enumerate(row):
                        cell = shape.table.cell(i,j)
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = RGBColor.from_string(INK if i==0 else ('E9F3F4' if i%2 else 'F4F7F8'))
                        cell.margin_left = Inches(.16)
                        cell.margin_top = Inches(.16)
                        cell.text = value
                        font(cell.text_frame.paragraphs[0].runs[0], 19, 'FFFFFF' if i==0 else INK, i==0)
            else:
                raise ValueError('Unsupported object type: '+obj['type'])
        text(slide, spec['id']+'/page', f'{idx:02d} / {len(data["slides"]):02d}', [11.45, 7.0, 1.1, .3], 11, 'FFFFFF' if spec['role']=='cover' else MUTED)
        slide.notes_slide.notes_text_frame.text = spec['notes']
    prs.save(output)

if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--source', required=True)
    parser.add_argument('--output', required=True)
    parser.add_argument('--init', action='store_true')
    args = parser.parse_args()
    path = Path(args.source)
    if args.init:
        if path.exists():
            raise SystemExit('Refusing to overwrite source; omit --init to rebuild.')
        path.write_text(json.dumps(source(), ensure_ascii=False, indent=2), encoding='utf-8')
    build(json.loads(path.read_text(encoding='utf-8')), args.output)
    print(args.output)
