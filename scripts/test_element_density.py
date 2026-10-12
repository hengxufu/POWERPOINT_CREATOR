"""Synthetic offline tests for the user's 40-object rule."""
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from audit_element_density import audit


class DensityTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.deck = Path(self.temp.name) / 'synthetic.pptx'
        self.prs = Presentation()
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])

    def tearDown(self):
        self.temp.cleanup()

    def add_text(self, shapes, index, text='Synthetic'):
        shape = shapes.add_textbox(Inches(index % 10 * .8), Inches(index // 10 * .5), Inches(.7), Inches(.4))
        shape.text = text
        return shape

    def result(self):
        self.prs.save(self.deck)
        return audit(self.deck)

    def test_39_requires_one_independent_web_asset(self):
        for index in range(39):
            self.add_text(self.slide.shapes, index)
        result = self.result()
        self.assertFalse(result['ready'])
        self.assertEqual(result['total_deficit'], 1)
        self.assertEqual(len(result['slides'][0]['suggested_web_asset_requests']), 1)
        self.assertEqual(result['request_status'], 'proposals_not_submitted')

    def test_40_passes_and_paragraphs_do_not_inflate_count(self):
        for index in range(40):
            shape = self.add_text(self.slide.shapes, index)
            shape.text_frame.add_paragraph().text = 'Extra paragraph'
        result = self.result()
        self.assertTrue(result['ready'])
        self.assertEqual(result['slides'][0]['element_count'], 40)

    def test_group_is_expanded_and_empty_hidden_offcanvas_are_excluded(self):
        group = self.slide.shapes.add_group_shape()
        self.add_text(group.shapes, 0)
        self.add_text(group.shapes, 1)
        self.add_text(self.slide.shapes, 2, '')
        outside = self.add_text(self.slide.shapes, 3)
        outside.left = Inches(30)
        hidden = self.add_text(self.slide.shapes, 4)
        hidden._element.xpath('.//p:cNvPr')[0].set('hidden', '1')
        self.assertEqual(self.result()['slides'][0]['element_count'], 2)

    def test_table_is_one_object_and_all_pages_are_checked(self):
        self.slide.shapes.add_table(5, 5, Inches(1), Inches(1), Inches(5), Inches(3))
        self.prs.slides.add_slide(self.prs.slide_layouts[6])
        result = self.result()
        self.assertEqual([page['element_count'] for page in result['slides']], [1, 0])
        self.assertEqual(result['total_deficit'], 79)


if __name__ == '__main__':
    unittest.main()
