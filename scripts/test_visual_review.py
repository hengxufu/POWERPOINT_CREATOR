"""SYNTHETIC board and review records only; never production visual evidence."""
import copy
import tempfile
import unittest
from pathlib import Path
from PIL import Image
from pptx import Presentation
from material_handoff import write, load, digest
from visual_review import prepare, check, CRITERIA


def fixture(root):
    (root/'planning').mkdir(parents=True,exist_ok=True)
    (root/'renders').mkdir(exist_ok=True)
    (root/'build').mkdir(exist_ok=True)
    plan={'representative_slides':{'cover':'s1','technical':'s2','dense':'s3'},'slides':[
        {'slide_number':n,'slide_id':f's{n}','title':'SYNTHETIC '+str(n),'theme_variant':'cover' if n==1 else 'content',
         'takeaway':'Synthetic test only','layout_family':'figure-analysis','visual_rationale':'Synthetic layout',
         'reading_order':['claim'], 'content_blocks':[{'object_id':'claim','role':'primary','text':'Synthetic explanation',
                                                     'font_pt':22,'box_inches':[1,1.8,7,3]}]}
        for n in range(1,5)]}
    write(root/'planning/deck_plan.json',plan)
    deck=root/'build/synthetic.pptx'
    prs=Presentation()
    for _ in plan['slides']:prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(deck)
    renders=[]
    for n in range(1,5):
        path=root/f'renders/synthetic-{n}.png'
        Image.new('RGB',(32,18),(20*n,50,90)).save(path)
        renders.append({'slide_number':n,'render_file':str(path.relative_to(root)),'render_sha256':digest(path)})
    index=root/'renders/synthetic-index.json'
    write(index,{'pptx_sha256':digest(deck),'slides':renders})
    return plan,deck,index


def synthetic_pass(path):
    record=load(path)
    for page in record['pages']:
        page.update(status='pass',note='SYNTHETIC TEST ATTESTATION ONLY',checks={key:True for key in CRITERIA})
    write(path,record)


class VisualTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.root=Path(self.temp.name).resolve()
        self.plan,self.deck,self.index=fixture(self.root)

    def tearDown(self):self.temp.cleanup()

    def test_board_is_pending_and_repeat_preserves_observations(self):
        result=prepare(self.root,self.deck,self.index)
        path=Path(result['review_file'])
        self.assertFalse(check(self.root,self.plan,'representative')['ready'])
        self.assertEqual(len(load(path)['pages']),3)
        self.assertIn('data:image/png;base64,',Path(result['board_file']).read_text(encoding='utf-8'))
        synthetic_pass(path)
        before=path.read_bytes()
        prepare(self.root,self.deck,self.index)
        self.assertEqual(path.read_bytes(),before)
        self.assertTrue(check(self.root,self.plan,'representative')['ready'])

    def test_actual_render_hash_and_final_deck_binding(self):
        data=load(self.index)
        data['pptx_sha256']='wrong'
        write(self.index,data)
        with self.assertRaisesRegex(ValueError,'another PPTX'):prepare(self.root,self.deck,self.index)
        data['pptx_sha256']=digest(self.deck)
        data['slides'][0]['render_sha256']='wrong'
        write(self.index,data)
        with self.assertRaisesRegex(ValueError,'hash mismatch'):prepare(self.root,self.deck,self.index)

    def test_full_and_representative_are_separate_and_require_all_criteria(self):
        rep=prepare(self.root,self.deck,self.index)
        synthetic_pass(rep['review_file'])
        self.assertFalse(check(self.root,self.plan,'full',self.deck)['ready'])
        full=prepare(self.root,self.deck,self.index,'full')
        synthetic_pass(full['review_file'])
        self.assertTrue(check(self.root,self.plan,'full',self.deck)['ready'])
        data=load(full['review_file']); data['pages'][3]['checks']['evidence']=False
        write(full['review_file'],data)
        self.assertFalse(check(self.root,self.plan,'full',self.deck)['ready'])

    def test_representative_survives_other_page_edits_but_design_changes_invalidate(self):
        result=prepare(self.root,self.deck,self.index)
        synthetic_pass(result['review_file'])
        changed=copy.deepcopy(self.plan)
        changed['slides'][3]['title']='Different non-representative page'
        self.assertTrue(check(self.root,changed,'representative')['ready'])
        changed['design_system']={'color_roles':{'text':'#112233'}}
        self.assertFalse(check(self.root,changed,'representative')['ready'])

    def test_snapshot_tampering_and_wrong_final_deck_fail(self):
        result=prepare(self.root,self.deck,self.index,'full')
        synthetic_pass(result['review_file'])
        other=self.root/'build/other.pptx'
        prs=Presentation(self.deck); prs.slides[0].shapes.add_textbox(0,0,100000,100000).text='SYNTHETIC change';prs.save(other)
        self.assertFalse(check(self.root,self.plan,'full',other)['ready'])
        path=Path(result['review_file']).parent/'slide-1.png'
        Image.new('RGB',(32,18),'red').save(path)
        self.assertFalse(check(self.root,self.plan,'full',self.deck)['ready'])

    def test_duplicate_missing_and_escaped_render_paths_fail(self):
        original=load(self.index)
        variants=[{**original,'slides':original['slides']*2}, {**original,'slides':original['slides'][1:]}]
        escaped=copy.deepcopy(original); escaped['slides'][0]['render_file']='../outside.png';variants.append(escaped)
        for data in variants:
            write(self.index,data)
            with self.assertRaises(ValueError):prepare(self.root,self.deck,self.index)


if __name__=='__main__':unittest.main()
