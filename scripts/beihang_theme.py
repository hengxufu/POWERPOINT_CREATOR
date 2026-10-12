"""Apply and audit a source-grounded, recurring Beihang identity layer."""
import argparse
import hashlib
import json
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_COLOR_TYPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

PACK = Path(__file__).resolve().parent.parent / 'assets/beihang'
PROFILES = {'beihang-blue': '005BAB', 'beihang-green': '357452', 'beihang-ginkgo': '9F9C2B'}
PREFIX = 'beihang-theme/'
ROLES = ('identity', 'university', 'section', 'page')


def load(path):
    return json.loads(Path(path).read_text(encoding='utf-8-sig'))


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def contract(profile='beihang-blue'):
    if profile not in PROFILES:
        raise ValueError('Unknown Beihang theme profile')
    manifest = load(PACK / 'manifest.json')
    return {'theme_profile': profile, 'university': '北京航空航天大学',
            'primary_color': PROFILES[profile], 'recurring_elements': list(ROLES),
            'source_assets': manifest['assets'], 'source_manifest_sha256': sha(PACK / 'manifest.json'),
            'reserved_regions_fraction': {'identity': [.79, .025, .185, .075], 'footer': [0, .925, 1, .075]},
            'brand_policy': 'Reuse original identity bytes; never generate or distort university logos; source colors are not verified official VI.'}


def sections(plan, total):
    slides = plan.get('slides', [])
    if [item.get('slide_number') for item in slides] != list(range(1, total + 1)):
        raise ValueError('Theme plan must cover every slide in actual display order')
    result = []
    for slide in slides:
        section = slide.get('section')
        if not isinstance(section, str) or not section.strip() or len(section) > 24:
            raise ValueError('Every page requires a concrete section label of 1..24 characters')
        result.append(section.strip())
    return result


def geometry(prs):
    w, h = prs.slide_width, prs.slide_height
    # Never stretch the original combination mark; contain it in the reserved box.
    with Image.open(PACK / 'wordmark-blue.png') as image:
        ratio = image.width / image.height
    iw = min(w * .185, h * .075 * ratio)
    ih = iw / ratio
    return {'identity': [int(w * .975 - iw), int(h * .025), int(iw), int(ih)],
            'university': [int(w * .04), int(h * .942), int(w * .29), int(h * .038)],
            'section': [int(w * .38), int(h * .942), int(w * .44), int(h * .038)],
            'page': [int(w * .87), int(h * .942), int(w * .1), int(h * .038)]}


def add_text(slide, role, text, box, color, scale):
    shape = slide.shapes.add_textbox(*box)
    shape.name = PREFIX + role
    frame = shape.text_frame
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    run = frame.paragraphs[0].add_run()
    run.text = text
    run.font.name = 'Times New Roman'
    run.font.size = Pt(10.5 * scale)
    run.font.color.rgb = RGBColor.from_string(color)
    east = OxmlElement('a:ea')
    east.set('typeface', '微软雅黑')
    run._r.get_or_add_rPr().append(east)
    return shape


def overlaps(a, b):
    return min(a[0]+a[2], b[0]+b[2]) > max(a[0], b[0]) and min(a[1]+a[3], b[1]+b[3]) > max(a[1], b[1])


def collisions(prs):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    findings = []
    theme_boxes = geometry(prs)
    def visit(shapes, number):
        for shape in shapes:
            if shape.name.startswith(PREFIX):
                continue
            metadata = shape._element.xpath('.//p:cNvPr')
            if metadata and metadata[0].get('hidden') in ('1', 'true'):
                continue
            box = [shape.left, shape.top, shape.width, shape.height]
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                # Group coordinates need transforms; use its outer bounds conservatively.
                pass
            elif shape.has_text_frame and not shape.text.strip() and not hasattr(shape, 'image') and not shape.has_chart and not shape.has_table:
                if shape.width * shape.height < prs.slide_width * prs.slide_height * .7:
                    continue
            # Large background images/fills are intentional layers, not foreground conflicts.
            if shape.width * shape.height >= prs.slide_width * prs.slide_height * .7:
                continue
            roles = [role for role, reserved in theme_boxes.items() if overlaps(box, reserved)]
            if roles:
                findings.append({'slide_number': number, 'object': shape.name, 'overlaps_theme': roles})
    for number, slide in enumerate(prs.slides, 1):
        visit(slide.shapes, number)
    return findings


def audit(pptx, plan, profile='beihang-blue'):
    theme = contract(profile)
    prs = Presentation(pptx)
    labels = sections(plan, len(prs.slides))
    expected = geometry(prs)
    problems, pages = [], []
    original_hash = theme['source_assets']['wordmark-blue.png']['sha256']
    if sha(PACK / 'wordmark-blue.png') != original_hash:
        problems.append('Pack identity hash differs from source manifest')
    planned_profile = plan.get('design_system', {}).get('theme_profile')
    if planned_profile != profile:
        problems.append('Plan design_system.theme_profile does not match theme contract')
    for number, slide in enumerate(prs.slides, 1):
        page_problems = []
        expected_text = {'university': '北京航空航天大学', 'section': labels[number-1], 'page': f'{number:02d} / {len(prs.slides):02d}'}
        for role in ROLES:
            shapes = [shape for shape in slide.shapes if shape.name == PREFIX + role]
            if len(shapes) != 1:
                page_problems.append(f'{role}: requires exactly one visible top-level object')
                continue
            shape = shapes[0]
            metadata = shape._element.xpath('.//p:cNvPr')
            if metadata and metadata[0].get('hidden') in ('1', 'true'):
                page_problems.append(f'{role}: hidden object')
            actual = [shape.left, shape.top, shape.width, shape.height]
            if any(abs(a-b) > 9144 for a, b in zip(actual, expected[role])) or shape.rotation != 0:
                page_problems.append(f'{role}: changed position, dimensions or rotation')
            if role == 'identity':
                if not hasattr(shape, 'image') or hashlib.sha256(shape.image.blob).hexdigest() != original_hash:
                    page_problems.append('identity: not original source image')
                if hasattr(shape, 'crop_left') and any(abs(v) > 1e-6 for v in (shape.crop_left, shape.crop_right, shape.crop_top, shape.crop_bottom)):
                    page_problems.append('identity: cropped mark')
                if abs(shape.width / max(1, shape.height) - expected[role][2] / expected[role][3]) > .0005:
                    page_problems.append('identity: distorted aspect ratio')
                if shape._element.xpath('.//a:alphaModFix | .//a:alphaBiLevel | .//a:duotone | .//a:grayscl | .//a:lum | .//a:clrChange'):
                    page_problems.append('identity: modified image effects')
            else:
                if not shape.has_text_frame or shape.text != expected_text[role]:
                    page_problems.append(f'{role}: wrong native text')
                elif any(run.font.color.type != MSO_COLOR_TYPE.RGB or str(run.font.color.rgb) != PROFILES[profile]
                         for paragraph in shape.text_frame.paragraphs for run in paragraph.runs):
                    page_problems.append(f'{role}: inconsistent theme color')
        pages.append({'slide_number': number, 'problems': page_problems})
        problems.extend(f'Page {number}: {item}' for item in page_problems)
    conflicts = collisions(prs)
    problems.extend(f"Page {item['slide_number']}: {item['object']} overlaps theme reserved area" for item in conflicts)
    return {'schema': 'beihang_theme_audit_v1', 'pptx_sha256': sha(pptx), 'theme_profile': profile,
            'source_manifest_sha256': theme['source_manifest_sha256'], 'ready': bool(pages) and not problems,
            'problems': problems, 'slides': pages, 'collisions': conflicts,
            'limits': 'Object/hash/layout check only; inspect full-deck renders for occlusion, readability and aesthetic consistency.'}


def apply(pptx, plan, out, profile='beihang-blue'):
    pptx, out = Path(pptx), Path(out)
    if pptx.resolve() == out.resolve() or out.exists():
        raise ValueError('Output must be a new PPTX; preserve original input')
    theme = contract(profile)
    if sha(PACK / 'wordmark-blue.png') != theme['source_assets']['wordmark-blue.png']['sha256']:
        raise ValueError('Source identity hash mismatch')
    prs = Presentation(pptx)
    labels = sections(plan, len(prs.slides))
    if plan.get('design_system', {}).get('theme_profile') != profile:
        raise ValueError('Plan theme_profile must match chosen profile')
    boxes = geometry(prs)
    for number, slide in enumerate(prs.slides, 1):
        for shape in list(slide.shapes):
            if shape.name.startswith(PREFIX):
                if shape.name not in {PREFIX + role for role in ROLES}:
                    raise ValueError('Unknown object in reserved theme namespace; refusing to remove it')
                slide.shapes._spTree.remove(shape._element)
        shape = slide.shapes.add_picture(str(PACK / 'wordmark-blue.png'), *boxes['identity'])
        shape.name = PREFIX + 'identity'
        shape._element.xpath('.//p:cNvPr')[0].set('descr', '北京航空航天大学校徽与中英校名组合；用户本地模板原始素材')
        for role, text in {'university': '北京航空航天大学', 'section': labels[number-1], 'page': f'{number:02d} / {len(prs.slides):02d}'}.items():
            add_text(slide, role, text, boxes[role], PROFILES[profile], prs.slide_height / 914400 / 7.5)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return audit(out, plan, profile)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('command', choices=['apply', 'audit'])
    parser.add_argument('--pptx', required=True, type=Path)
    parser.add_argument('--plan', required=True, type=Path)
    parser.add_argument('--profile', choices=PROFILES, default='beihang-blue')
    parser.add_argument('--out', required=True, type=Path)
    parser.add_argument('--report', type=Path)
    args = parser.parse_args()
    protected = {args.pptx.resolve(), args.plan.resolve()}
    if args.out.resolve() in protected or (args.report and args.report.resolve() in protected | {args.out.resolve()}):
        parser.error('Outputs must not overwrite inputs or each other')
    if args.command == 'apply':
        if not args.report:
            parser.error('apply requires --report')
        result = apply(args.pptx, load(args.plan), args.out, args.profile)
        report = args.report
    else:
        result = audit(args.pptx, load(args.plan), args.profile)
        report = args.out
    report.parent.mkdir(parents=True, exist_ok=True)
    report.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'ready': result['ready'], 'problems': result['problems']}, ensure_ascii=False))
    raise SystemExit(0 if result['ready'] else 1)
