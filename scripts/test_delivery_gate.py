"""Synthetic offline regression checks; never represent fixtures as actual web output."""
import copy
import json
import tempfile
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from delivery_gate import check, digest


def main():
    passed = []
    with tempfile.TemporaryDirectory() as temporary:
        root = Path(temporary)
        for folder in ('planning', 'qa', 'assets/job', 'renders'):
            (root / folder).mkdir(parents=True)
        image = root / 'assets/job/icon.png'
        Image.new('RGB', (64, 64), 'blue').save(image)
        render = root / 'renders/slide-01.png'
        Image.new('RGB', (128, 72), 'white').save(render)
        deck = root / 'deck.pptx'
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(str(image), Inches(1), Inches(1)).name = 's1/icon'
        for index in range(39):
            text = slide.shapes.add_textbox(Inches((index % 10) * .8), Inches(2 + (index // 10) * .5), Inches(.7), Inches(.4))
            text.text = f'Test {index}'
        prs.save(deck)
        plan = {'slides': [{'slide_number': 1, 'title': 'Synthetic fixture', 'takeaway': 'Test evidence reconciliation',
                            'layout_family': 'figure-analysis', 'visual_rationale': 'Synthetic test only',
                            'content_blocks': [{'text': 'A specific fixture explanation'}],
                            'asset_requests': [{'request_id': 'r1', 'decision': 'generate_web'}]}]}
        review = {'pptx_sha256': digest(deck), 'unresolved_issues': [],
                  'web_planning': {'session_url': 'https://chatgpt.com/c/SYNTHETIC-FIXTURE', 'response_file': 'planning/response.txt'},
                  'slides': [{'slide_number': 1, 'render_file': 'renders/slide-01.png', 'render_sha256': digest(render),
                              'rendered_from_sha256': digest(deck),
                              **{key: {'status': 'pass', 'note': 'Synthetic test attestation only'} for key in ('content_review', 'visual_review', 'evidence_review')}}]}
        job = {'transport': 'chatgpt_web', 'status': 'accepted', 'history': [{'status': 'submitted'}],
               'spec': {'slide_id': 's1', 'object_id': 'icon'}, 'sha256': digest(image),
               'visual_review': {'note': 'Synthetic fixture only'}, 'verified_engine': None}
        manifest = {'elements': [{'request_id': 'r1', 'slide_number': 1, 'job_dir': 'job'}]}
        (root / 'planning/response.txt').write_text('Synthetic test response; not an actual web exchange.', encoding='utf-8')
        from bridge_session import prepare
        connection = prepare(root, root)
        connection.update(local_profile_configured=True, browser_connected=True, image_tool_visible=True,
                          planning_model='GPT-6 (synthetic fixture)', image_model='Image-2.5 (synthetic fixture)',
                          observed_connection_id=connection['connection_id'],
                          evidence=[{'file': 'renders/slide-01.png', 'sha256': digest(render)}])
        (root / 'bridge-connect.json').write_text(json.dumps(connection), encoding='utf-8')

        def write(relative, data):
            (root / relative).write_text(json.dumps(data), encoding='utf-8')

        def reset():
            for relative, data in [('planning/deck_plan.json', plan), ('qa/delivery_review.json', review),
                                   ('assets/job/job.json', job), ('assets/asset_manifest.json', manifest)]:
                write(relative, data)

        reset()
        assert check(root, deck)['ready']
        assert check(root, deck)['warnings']
        passed.append('complete synthetic records reconcile; unverified engine remains disclosed')
        mutations = [
            ('assets/asset_manifest.json', {'elements': []}, 'empty manifest cannot skip planned image'),
            ('assets/job/job.json', {**job, 'status': 'prepared'}, 'prepared prompt is not completed generation'),
            ('assets/job/job.json', {**job, 'history': []}, 'accepted label without submission fails'),
            ('assets/job/job.json', {**job, 'sha256': 'wrong'}, 'wrong picture hash fails'),
            ('assets/asset_manifest.json', {'elements': [{**manifest['elements'][0], 'slide_number': 2}]}, 'wrong destination fails'),
            ('qa/delivery_review.json', {**review, 'pptx_sha256': 'old'}, 'stale deck review fails'),
            ('qa/delivery_review.json', {**review, 'slides': []}, 'missing page review fails'),
            ('qa/delivery_review.json', {**review, 'web_planning': {}}, 'missing web response fails'),
        ]
        for relative, data, name in mutations:
            reset()
            write(relative, data)
            assert not check(root, deck)['ready'], name
            passed.append(name)
        reset()
        bad = copy.deepcopy(review)
        bad['slides'][0]['render_sha256'] = 'old'
        write('qa/delivery_review.json', bad)
        assert not check(root, deck)['ready']
        passed.append('stale render fails')
        reset()
        bad = copy.deepcopy(plan)
        bad['slides'][0]['content_blocks'] = [{'heading': 'Background'}]
        write('planning/deck_plan.json', bad)
        assert not check(root, deck)['ready']
        passed.append('heading-only content fails')
        reset()
        # Theme-enabled synthetic workspace must reconcile actual source images, not names alone.
        from beihang_theme import apply as apply_theme, PREFIX
        theme_plan = copy.deepcopy(plan)
        theme_plan['design_system'] = {'theme_profile': 'beihang-blue'}
        theme_plan['slides'][0]['section'] = '合成测试'
        theme_plan['slides'][0]['slide_id'] = 's1'
        themed_deck = root / 'themed.pptx'
        theme_result = apply_theme(deck, theme_plan, themed_deck)
        assert theme_result['ready'], theme_result['problems']
        write('composition_contract.json', {'theme_profile': 'beihang-blue'})
        write('planning/deck_plan.json', theme_plan)
        themed_review = copy.deepcopy(review)
        themed_review['pptx_sha256'] = digest(themed_deck)
        themed_review['slides'][0]['rendered_from_sha256'] = digest(themed_deck)
        write('qa/delivery_review.json', themed_review)
        assert check(root, themed_deck)['ready']
        passed.append('enabled Beihang theme reconciles original source image and recurring objects')
        damaged = Presentation(themed_deck)
        identity = next(s for s in damaged.slides[0].shapes if s.name == PREFIX + 'identity')
        damaged.slides[0].shapes._spTree.remove(identity._element)
        damaged.save(themed_deck)
        themed_review['pptx_sha256'] = digest(themed_deck)
        themed_review['slides'][0]['rendered_from_sha256'] = digest(themed_deck)
        write('qa/delivery_review.json', themed_review)
        result = check(root, themed_deck)
        assert not result['ready'] and any('Theme:' in p for p in result['problems'])
        passed.append('missing recurring identity fails even when final deck hash and density reconcile')
    for name in passed:
        print('PASS:', name)
    print(f'{len(passed)} offline checks passed')


if __name__ == '__main__':
    main()
