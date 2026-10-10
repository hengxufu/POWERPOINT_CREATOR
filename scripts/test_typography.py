"""Regression checks: missing fonts, charts, field text, math and package fidelity."""
import importlib.util
import tempfile
import zipfile
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
from lxml import etree as E

HERE = Path(__file__).resolve().parent
spec = importlib.util.spec_from_file_location('fontcheck', HERE / 'typography_check.py')
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)


def main():
    checks = []
    with tempfile.TemporaryDirectory() as directory:
        root = Path(directory)
        source, fixed = root / 'source.pptx', root / 'fixed.pptx'
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = '科研汇报 Results 2026'
        prs.save(source)
        result = module.audit(source)
        assert not result['passed_explicit_font_check']
        assert {item['script'] for item in result['missing_run_fonts'] if item['part'] == 'ppt/slides/slide1.xml'} == {'ea', 'latin'}
        checks.append('undeclared mixed-script run fails audit')
        # Inject a chart text field and an equation; preserve AlternateContent and non-XML data.
        with zipfile.ZipFile(source, 'a') as archive:
            archive.writestr('ppt/charts/chart99.xml', f'''<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" xmlns:a="{module.A}"><c:txPr><a:p><a:fld id="field"><a:rPr b="1"/><a:t>图表 Chart</a:t></a:fld></a:p></c:txPr></c:chartSpace>''')
            archive.writestr('ppt/slides/equation-fixture.xml', f'''<p:sld xmlns:p="{module.P}" xmlns:a="{module.A}" xmlns:m="{module.M}" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"><mc:AlternateContent><mc:Choice Requires="m"><m:oMath><a:r><a:rPr><a:latin typeface="Cambria Math"/></a:rPr><a:t>x</a:t></a:r></m:oMath></mc:Choice><mc:Fallback><p:pic/></mc:Fallback></mc:AlternateContent></p:sld>''')
            archive.writestr('ppt/media/protected.bin', b'unchanged-image-and-workbook-fixture')
        module.normalize(source, fixed)
        result = module.audit(fixed)
        assert result['passed_explicit_font_check'], result
        assert result['counts']['office_equations'] == 1
        checks.append('normalization passes including chart field text')
        with zipfile.ZipFile(source) as before, zipfile.ZipFile(fixed) as after:
            assert before.namelist() == after.namelist()
            assert all(before.read(name) == after.read(name) for name in before.namelist() if not module.relevant(name))
            checks.append('all relationships media and other untouched parts preserved byte-for-byte')
            old = module.parse(before.read('ppt/slides/equation-fixture.xml'))
            new = module.parse(after.read('ppt/slides/equation-fixture.xml'))
            assert E.tostring(old, method='c14n') == E.tostring(new, method='c14n')
            checks.append('OMML and AlternateContent structure preserved')
            chart = module.parse(after.read('ppt/charts/chart99.xml'))
            properties = next(chart.iter(f'{{{module.A}}}rPr'))
            assert properties.get('b') == '1'
            checks.append('bold formatting preserved')
        assert Presentation(fixed).slides[0].shapes[0].text == '科研汇报 Results 2026'
        checks.append('deck reopens with original editable text')
        for output in (source, fixed):
            try:
                module.normalize(source, output)
            except ValueError:
                pass
            else:
                raise AssertionError('overwrite accepted')
        checks.append('source and existing outputs cannot be overwritten')
    for check in checks:
        print('PASS:', check)
    print(f'{len(checks)} regression checks passed')


if __name__ == '__main__':
    main()
